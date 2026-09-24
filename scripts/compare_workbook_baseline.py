"""Read-only workbook comparison against a frozen, source-reviewed JSON reference.

This development tool does not import the production extraction/parser pipeline.
It preserves source row multiplicities; equal rows are not automatically duplicates.
Exit 0: no detected differences (not certification); 1: differences; 2: input error.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict, deque
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook


def normalized(text):
    # Notes stay in source evidence; ignore their spacing when matching labels.
    text = re.sub(r"\(\s*(?:Note\s+)?\d+\s*\)", "", str(text or ""), flags=re.I)
    return re.sub(r"\s+", " ", text).strip().rstrip(":").casefold()


def amount(value):
    if value is None or value in ("", "\u2014", "\u2013", "-"):
        return None
    if isinstance(value, bool):
        raise ValueError("Boolean is not a financial amount")
    text = str(value).replace(",", "").strip()
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Not an amount: {value!r}") from exc
    if not result.is_finite():
        raise ValueError("Nonfinite amount")
    return result


def is_amount(value):
    if value is None or value == "":
        return False
    try:
        amount(value)
        return True
    except ValueError:
        return False


def signature(row):
    return normalized(row["label"]), tuple(amount(v) for v in row["values"])


def compare_rows(expected, actual):
    """Match complete rows before pairing changed rows with the same source label."""
    available = defaultdict(deque)
    for index, row in enumerate(expected):
        available[signature(row)].append(index)
    originals = set(available)
    matched, unmatched, duplicates = [], [], []
    for row in actual:
        key = signature(row)
        if available[key]:
            matched.append((available[key].popleft(), row))
        elif key in originals:
            duplicates.append(row)
        else:
            unmatched.append(row)
    remaining = sorted(i for queue in available.values() for i in queue)
    changed, unstructured, unexpected = [], [], []
    for row in unmatched:
        candidates = [i for i in remaining if normalized(expected[i]["label"]) == normalized(row["label"])]
        if not candidates:
            unexpected.append(row)
            continue
        # Prefer the most similar row when a label repeats across source sections.
        i = max(candidates, key=lambda index: sum(a == b for a, b in zip(
            signature(expected[index])[1], signature(row)[1])))
        remaining.remove(i)
        pair = {"source": expected[i], "actual": row}
        if len(expected[i]["values"]) != len(row["values"]):
            unstructured.append(pair)
        else:
            pair["changed_fields"] = [j for j, (a, b) in enumerate(zip(
                signature(expected[i])[1], signature(row)[1])) if a != b]
            changed.append(pair)
    order = [i for i, _ in matched]
    return dict(matched_rows=len(matched),
                matched_numeric_values=sum(v is not None for i, _ in matched for v in expected[i]["values"]),
                matches=[{"source_id": expected[i].get("id", str(i)), "workbook_row": row.get("row"),
                          "columns": row.get("columns", [])} for i, row in matched],
                missing_rows=[expected[i] for i in remaining],
                duplicated_rows=duplicates, changed_rows=changed, unstructured_rows=unstructured,
                unexpected_rows=unexpected, order_changed=order != sorted(order))


def read_statement(sheet, reference, *, label_column=1):
    source_labels = {normalized(r["label"]): r["label"] for r in reference["rows"] if r["label"]}
    width = len(reference["fields"])
    provisional, unexpected = [], []
    for cells in sheet.iter_rows(min_row=5, min_col=label_column):
        excel_row = cells[0].row
        label = str(cells[0].value or "")
        tokens = [(c.column, c.value) for c in cells[1:] if is_amount(c.value)]
        # A prose footnote can mention a source label and contain dates/note
        # numbers. With no separate amounts or exact source-row match it is
        # prose, regardless of the worksheet's alignment or cell formatting.
        if (not tokens and normalized(label) not in source_labels
                and re.match(r'^\(\d+\)\s+(?:Includes\b|See\s+Notes?\b)', label, re.I)):
            continue
        if not tokens and not label:
            continue
        # Multi-line reconciliation labels may surround their amount row.
        exact = normalized(label)
        matches = [(exact, label)] if exact in source_labels else [
            (key, label) for key in source_labels if key in exact]
        # Only expand an isolated note marker, not an ordinary complete label,
        # a year-header row, or an unlabeled subtotal next to a named investment.
        if not matches and tokens and re.fullmatch(r"\(\d+\)", label.strip()):
            text = " ".join(str(sheet.cell(r, label_column).value or "")
                            for r in range(max(5, excel_row - 1), min(sheet.max_row, excel_row + 1) + 1))
            matches = [(key, text) for key in source_labels if key in normalized(text)]
        if matches:
            key, matched_text = max(matches, key=lambda pair: (len(pair[0]), -len(pair[1])))
            if not tokens:
                # Keep all-dash/blank source rows, but not a section sharing a data label.
                candidates = [r for r in reference["rows"] if normalized(r["label"]) == key]
                embedded_data = exact != key and bool(re.search(r"\d|\u2014", label))
                if not embedded_data and not any(all(v is None for v in r["values"]) for r in candidates):
                    continue
            source_label = source_labels[key]
        elif tokens and any(not r["label"] for r in reference["rows"]) and not re.search(r"[A-Za-z]", label):
            source_label, matched_text = "", label
        else:
            if tokens and re.search(r"[A-Za-z]", label) and not any(x in normalized(label) for x in (
                "in thousands", "months ended", "investment type", "years ended")):
                unexpected.append({"row": excel_row, "label": label, "values": [v for _, v in tokens]})
            continue
        if reference["id"] == "schedule":
            tokens = tokens[-width:]  # source monetary fields; rate/spread text is separately retained
        provisional.append(dict(row=excel_row, label=source_label, display_label=matched_text,
                                tokens=tokens, polluted_label=normalized(matched_text) != normalized(source_label)))
    columns = sorted({c for row in provisional for c, _ in row["tokens"]})
    actual = []
    for row in provisional:
        if len(columns) == width:
            positions = columns
            values = [sheet.cell(row["row"], c).value for c in columns]
            # Unsupported nonnumeric values must remain visible as a layout issue.
            if any(v is not None and not is_amount(v) for v in values):
                positions, values = zip(*row["tokens"]) if row["tokens"] else ([], [])
        else:
            positions, values = zip(*row["tokens"]) if row["tokens"] else ([], [])
        actual.append({k: v for k, v in row.items() if k != "tokens"} |
                      {"values": list(values), "columns": list(positions)})
    data_rows = {r["row"] for r in actual}
    header_text = normalized(" ".join(str(c.value or "") for row in sheet if row[0].row not in data_rows for c in row))
    variants = Counter(tuple(r["columns"]) for r in actual if len(r["values"]) == width)
    total_rows = [r for r in actual if normalized(r["label"]).startswith(("total ", "net cash ", "net increase "))]
    return actual, dict(
        missing_headings=[h for h in reference["headings"] if normalized(h) not in header_text],
        column_patterns=[{"columns": list(k), "rows": v} for k, v in variants.items()],
        inconsistent_columns=len(variants) > 1,
        labels_with_extra_text=[r["row"] for r in actual if r["polluted_label"]],
        totals_without_bold=[r["row"] for r in total_rows if not sheet.cell(r["row"], label_column).font.bold],
        unexpected_numeric_rows=unexpected,
        replacement_characters=[c.coordinate for row in sheet for c in row
                                if isinstance(c.value, str) and "\ufffd" in c.value],
    )


def inspect_workbook(path, reference):
    book = load_workbook(path, data_only=False)
    try:
        results = []
        for section in reference["statement_sections"]:
            # Support both frozen legacy workbooks and exports with three blank
            # top rows and two blank left columns. Match the full title.
            sheets = [(s, c) for s in book for c in (1, 3)
                      if any(normalized(s.cell(r, c).value) == normalized(section["title"])
                             for r in (2, 5))]
            base = dict(id=section["id"], title=section["title"], source_pages=section["pages"],
                        expected_rows=len(section["rows"]), expected_numeric_values=section["numeric_count"])
            if len(sheets) != 1:
                results.append(base | dict(sheet=None, error="Statement missing" if not sheets else "Multiple matching sheets",
                                           matched_rows=0, missing_rows=section["rows"]))
                continue
            sheet, label_column = sheets[0]
            rows, layout = read_statement(sheet, section, label_column=label_column)
            results.append(base | dict(sheet=sheet.title, actual_rows=len(rows), layout=layout,
                                      presentation={"rows": sheet.max_row, "columns": sheet.max_column,
                                                    "merged_ranges": sorted(str(r) for r in sheet.merged_cells.ranges),
                                                    "column_widths": {k: v.width for k, v in sheet.column_dimensions.items()}}) |
                           compare_rows(section["rows"], rows))
        inventory = [{"sheet": s.title, "rows": s.max_row, "columns": s.max_column,
                      "tables": sum(any(str(c.value).startswith("Table:") for c in row[:3]) for row in s)} for s in book]
        return dict(file=str(path.resolve()), sha256=file_hash(path), statements=results, inventory=inventory)
    finally:
        book.close()


def file_hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify_baseline(folder):
    manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
    for entry in manifest["files"].values():
        path = (folder / entry["file"]).resolve()
        if not path.is_relative_to(folder.resolve()) or file_hash(path) != entry["sha256"]:
            raise ValueError(f"Baseline integrity check failed: {entry['file']}")
    if "expected.json" not in manifest["files"]:
        raise ValueError("Reference JSON is not registered in the baseline manifest")
    reference = json.loads((folder / "expected.json").read_text(encoding="utf-8"))
    if reference["source_sha256"] != manifest["files"]["source"]["sha256"]:
        raise ValueError("Reference belongs to a different source PDF")
    return manifest, reference


def has_differences(result):
    for s in result["statements"]:
        if s.get("error") or s.get("order_changed"):
            return True
        if any(s.get(key) for key in ["missing_rows", "duplicated_rows", "changed_rows", "unstructured_rows", "unexpected_rows"]):
            return True
        layout = s.get("layout", {})
        if any(layout.get(key) for key in ["missing_headings", "inconsistent_columns", "labels_with_extra_text",
                                          "unexpected_numeric_rows", "replacement_characters", "totals_without_bold"]):
            return True
    return False


def raw_blocks(sheet):
    tables, current = {}, None
    first_column = min((c.column for row in sheet for c in row if c.value is not None), default=1)
    for row in sheet.iter_rows(min_col=first_column):
        values = [c.value for c in row]
        if str(values[0]).startswith("Table:"):
            current = values[0][7:]
            tables[current] = []
        elif current:
            while values and values[-1] is None:
                values.pop()
            if values:
                tables[current].append(values)
    return tables


def compare_raw(old_path, new_path):
    old, new = load_workbook(old_path), load_workbook(new_path)
    try:
        output = {}
        for old_name, new_name in [("Raw Camelot Lattice", "Lattice"), ("Raw Camelot Stream", "Stream"), ("Raw PDFPlumber", "PDFPlumber")]:
            if old_name not in old or new_name not in new:
                continue
            a, b = raw_blocks(old[old_name]), raw_blocks(new[new_name])
            output[new_name] = {"old_tables": len(a), "new_tables": len(b),
                                "removed_tables": sorted(a.keys() - b.keys()), "added_tables": sorted(b.keys() - a.keys()),
                                "changed_tables": sorted(k for k in a.keys() & b.keys() if a[k] != b[k])}
        return output
    finally:
        old.close()
        new.close()


def markdown(report):
    lines = ["# Workbook repair baseline", "", "Reference: the frozen PDF, physical pages 3-9. Old output is not the correctness oracle.", "",
             "Dashes and blank numeric cells compare as missing, never as zero. Their original tokens remain in the reference and workbooks.", "",
             "A matched row means the selected numeric fields match in source order within that row. It does not certify headings, physical column placement, rates, or footnotes.", ""]
    for name, result in report["workbooks"].items():
        lines += [f"## {name}", "", "| Statement | Source rows | Matched rows | Missing rows | Excess duplicates | Changed rows | Unstructured rows |", "|---|---:|---:|---:|---:|---:|---:|"]
        for s in result["statements"]:
            lines.append(f"| {s['id']} | {s['expected_rows']} | {s['matched_rows']} | {len(s['missing_rows'])} | {len(s.get('duplicated_rows', []))} | {len(s.get('changed_rows', []))} | {len(s.get('unstructured_rows', []))} |")
        lines.append("")
        for s in result["statements"]:
            layout = s.get("layout", {})
            if s.get("error"):
                lines.append(f"- {s['id']}: {s['error']}.")
            if layout.get("missing_headings"):
                lines.append(f"- {s['id']} missing headings: " + "; ".join(layout["missing_headings"]) + ".")
            if layout.get("inconsistent_columns"):
                lines.append(f"- {s['id']}: values shift between physical column patterns {layout['column_patterns']}.")
            if layout.get("labels_with_extra_text"):
                lines.append(f"- {s['id']}: {len(layout['labels_with_extra_text'])} row labels contain additional text or embedded data.")
            if layout.get("totals_without_bold"):
                lines.append(f"- {s['id']}: {len(layout['totals_without_bold'])} total rows lack bold emphasis.")
        lines.append("")
    lines += ["## Limits", ""] + ["- " + item for item in report["limitations"]]
    lines += ["", "Full row IDs, PDF line references, workbook row numbers, mismatches, raw-table comparisons and file hashes are in comparison.json."]
    return "\n".join(lines) + "\n"


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("baseline", type=Path)
    parser.add_argument("--candidate", action="append", default=[], metavar="NAME=PATH")
    parser.add_argument("--report-dir", type=Path, required=True, help="A new directory; existing reports are never overwritten")
    args = parser.parse_args(argv)
    try:
        manifest, reference = verify_baseline(args.baseline)
        paths = {name: args.baseline / manifest["files"][name]["file"] for name in ["old", "new", "corrected"]}
        for candidate in args.candidate:
            name, filename = candidate.split("=", 1)
            if not name or name in paths:
                raise ValueError("Candidate names must be nonempty and unique")
            paths[name] = Path(filename)
        if args.report_dir.exists():
            raise FileExistsError("Choose a new report directory; existing reports are preserved")
        targets = [c.split("=", 1)[0] for c in args.candidate] or list(paths)
        report = dict(schema_version=1, assessment_targets=targets, baseline=str(args.baseline.resolve()),
                      comparator_sha256=file_hash(Path(__file__)), manifest_sha256=file_hash(args.baseline / "manifest.json"),
                      limitations=reference["limitations"], workbooks={name: inspect_workbook(path, reference) for name, path in paths.items()},
                      raw_comparisons={name: compare_raw(paths["old"], path) for name, path in paths.items() if name != "old"})
        args.report_dir.mkdir(parents=True, exist_ok=False)
        (args.report_dir / "comparison.json").write_text(json.dumps(report, indent=2, ensure_ascii=True), encoding="utf-8")
        (args.report_dir / "comparison.md").write_text(markdown(report), encoding="utf-8")
        print(args.report_dir / "comparison.md")
        return 1 if any(has_differences(report["workbooks"][name]) for name in targets) else 0
    except (OSError, ValueError, KeyError) as exc:
        parser.exit(2, f"error: {exc}\n")


if __name__ == "__main__":
    raise SystemExit(main())
