"""Independently validate saved Apple balance-sheet cells against frozen evidence.

No production parser, normalization, recovery or audit code is imported.
Exit 0: bounded checks pass; 1: differences/unverified checks; 2: invalid inputs.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from decimal import Decimal
import json
import math
from pathlib import Path

from openpyxl import load_workbook

from scripts.compare_apple_balance_baseline import load_reference
from scripts.compare_workbook_baseline import has_differences, inspect_workbook, normalized, read_statement


def audit_saved_cells(sheet, section, controls):
    """Associate repeated labels by source occurrence, never by matching values."""
    actual, _ = read_statement(sheet, section, label_column=3)
    expected_by_label, actual_by_label = defaultdict(list), defaultdict(list)
    for row in section['rows']:
        expected_by_label[normalized(row['label'])].append(row)
    for row in actual:
        actual_by_label[normalized(row['label'])].append(row)
    by_source_line = {}
    for label, expected in expected_by_label.items():
        exported = actual_by_label[label]
        if len(expected) != len(exported):
            continue  # Missing/extra repeated labels cannot be paired safely.
        for source, target in zip(expected, exported):
            for line in source['source_lines']:
                if line in by_source_line:
                    raise ValueError(f'Ambiguous reference source line: {line}')
                by_source_line[line] = target

    first_data_row = min((row['row'] for row in actual), default=sheet.max_row + 1)
    period_columns, date_checks = {}, []
    for period in section['fields']:
        headers = [cell for row in sheet.iter_rows(max_row=first_data_row - 1) for cell in row
                   if cell.value == period and cell.data_type == 's']
        column = headers[0].column if len(headers) == 1 else None
        if column is not None:
            period_columns[period] = column
        date_checks.append(dict(check='Literal date above amount column', period=period,
                                status='PASS' if column else 'FAIL',
                                cells=[cell.coordinate for cell in headers]))

    def cell_value(source_line, period):
        row = by_source_line.get(source_line)
        c = period_columns.get(period)
        index = section['fields'].index(period)
        if row is None or c is None or len(row['columns']) != len(section['fields']) or row['columns'][index] != c:
            return None, None
        cell = sheet.cell(row['row'], c)
        value = cell.value
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            return None, cell.coordinate
        return Decimal(str(value)), cell.coordinate

    checks = []
    if not controls or len({(c['check'], c['period']) for c in controls}) != len(controls):
        raise ValueError('Missing or duplicated independent arithmetic controls')
    for control in controls:
        period = control['period']
        if period not in section['fields'] or not control['source_input_lines']:
            raise ValueError('Arithmetic control has no supported period/components')
        inputs = [cell_value(line, period) for line in control['source_input_lines']]
        reported, reported_cell = cell_value(control['total_source_line'], period)
        calculated = sum(value for value, _ in inputs) if all(value is not None for value, _ in inputs) else None
        diff = reported - calculated if reported is not None and calculated is not None else None
        checks.append(dict(check=control['check'], period=period,
                           status='NOT_TESTED' if diff is None else 'PASS' if diff == 0 else 'FAIL',
                           input_cells=[cell for _, cell in inputs], reported_cell=reported_cell,
                           calculated=str(calculated) if calculated is not None else None,
                           reported=str(reported) if reported is not None else None,
                           difference=str(diff) if diff is not None else None))
    return date_checks, checks


def validate(baseline, candidate):
    reference = load_reference(baseline, required_files=('source-controls.json',))
    controls = json.loads((baseline / 'source-controls.json').read_text(encoding='utf-8'))
    comparison = inspect_workbook(candidate, reference)
    sections = reference['statement_sections']
    if len(sections) != 1 or sections[0]['id'] != 'balance':
        raise ValueError('Expected one independently reviewed balance-sheet reference')
    name = comparison['statements'][0]['sheet']
    if name is None:
        return dict(status='FAIL', comparison=comparison, dates=[], arithmetic=[],
                    detail='A uniquely titled balance sheet was not found; arithmetic is unverified.')
    book = load_workbook(candidate, data_only=False)
    try:
        dates, arithmetic = audit_saved_cells(book[name], sections[0], controls)
    finally:
        book.close()
    passed = not has_differences(comparison) and all(c['status'] == 'PASS' for c in dates + arithmetic)
    return dict(status='PASS' if passed else 'FAIL', comparison=comparison, dates=dates, arithmetic=arithmetic,
                scope='Apple balance sheet only. Exact saved-cell arithmetic; other statements are outside this acceptance check.')


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('baseline', type=Path)
    parser.add_argument('candidate', type=Path)
    args = parser.parse_args(argv)
    try:
        result = validate(args.baseline, args.candidate)
    except (OSError, ValueError, KeyError) as exc:
        parser.exit(2, f'error: {exc}\n')
    print(json.dumps(result, indent=2, ensure_ascii=True))
    return 0 if result['status'] == 'PASS' else 1


if __name__ == '__main__':
    raise SystemExit(main())
