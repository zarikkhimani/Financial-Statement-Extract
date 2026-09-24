"""Native Tk integration checks for the view/state/controller boundary."""

import threading
import tkinter as tk
from types import SimpleNamespace

import pytest

from app import FinancialExtractorApp
from financial_statement_extract.ui.lifecycle import JobPhase
from financial_statement_extract.ui.theme import apply_theme
from models import ExtractionResult


@pytest.fixture
def make_workspace(monkeypatch, tk_runtime):
    workspaces = []
    dialogs = []
    callback_errors = []
    ui_thread = threading.get_ident()
    monkeypatch.setattr(tk_runtime, "report_callback_exception", lambda *exc: callback_errors.append(exc))

    def record_dialog(kind, title, message, **_kwargs):
        assert threading.get_ident() == ui_thread
        dialogs.append((kind, title, message))

    monkeypatch.setattr("app.messagebox.showinfo", lambda *args, **kw: record_dialog("info", *args, **kw))
    monkeypatch.setattr("app.messagebox.showerror", lambda *args, **kw: record_dialog("error", *args, **kw))
    monkeypatch.setattr("app.filedialog.askopenfilename", lambda **_kw: "")
    monkeypatch.setattr("app.filedialog.askdirectory", lambda **_kw: "")

    def make(runner, **kwargs):
        root = tk.Toplevel(tk_runtime)
        root.withdraw()
        apply_theme(root)
        app = FinancialExtractorApp(root, runner=runner, **kwargs)
        workspaces.append(app)
        root.update_idletasks()
        return app, dialogs

    yield make
    for app in workspaces:
        root = app.root
        app.progress.stop()
        # Tests manually poll the queue; remove outstanding timers before the
        # next view is created. No test leaves a desktop window open.
        for callback in root.tk.call("after", "info"):
            root.after_cancel(callback)
        root.destroy()
    assert not callback_errors, callback_errors


def unexpected_runner(*_args):
    raise AssertionError("This interaction should not start extraction")


def test_native_fields_and_drop_update_draft_without_starting_work(make_workspace, tmp_path):
    app, dialogs = make_workspace(unexpected_runner)
    assert app.root.title() == "Financial Statement Extract"
    assert app.lifecycle.phase is JobPhase.EMPTY
    assert app.pages.get() == "auto"
    assert app.state.pages == "auto"
    assert app._metadata() == {"client_name": "", "year": "", "period": "Auto", "audit_status": "Audited"}
    app.pdf_run_button.invoke()
    assert not dialogs  # Empty-state action opens the chooser; cancelling is quiet.
    assert app.pdf_run_button.cget("text") == "Choose filing"
    app._start_pdf_extraction()
    assert dialogs[-1] == ("error", "Missing filing", "Choose or drop a PDF or HTML filing first.")

    source = tmp_path / "Company 10-K.pdf"
    source.write_bytes(b"%PDF-1.4")
    # TkDND uses its own event substitutions; ordinary event_generate cannot
    # synthesize an OS drop. Verify registration and feed its payload to the
    # handler, which uses the real Tcl interpreter to split paths with spaces.
    assert app.root.dnd_bind("<<Drop>>")
    assert app.pages_entry.dnd_bind("<<Drop>>")
    assert app._on_file_drop(SimpleNamespace(data=f"{{{source}}}", action="copy")) == "copy"
    assert app.state.input_path == str(source)
    assert app.lifecycle.phase is JobPhase.READY
    assert app.state.output_dir == str(tmp_path)
    assert "Loaded filing" in app.pdf_status.get("1.0", "end")

    app._on_pages_focus_in()
    app.pages_entry.insert(0, "12,14-16")
    app._on_pages_focus_out()
    assert app.state.pages == "12,14-16"
    app.pages_entry.delete(0, "end")
    app._on_pages_focus_out()
    assert app.state.pages == "auto"
    assert app.pages.get() == "auto"

    app.client_name.set(" Example ")
    app.year.set(" 2026 ")
    app.period.set("Annual")
    app.audit_status.set("Reviewed")
    app.output_dir.set(str(tmp_path / "chosen-output"))
    html = tmp_path / "Next filing.html"
    html.write_text("<html></html>", encoding="utf-8")
    assert app._set_pdf_path(html)
    assert app.state.output_dir == str(tmp_path / "chosen-output")
    assert app.state.to_request().metadata == {
        "client_name": "Example", "year": "2026", "period": "Annual", "audit_status": "Reviewed",
    }
    assert app.controller.worker is None


@pytest.mark.parametrize("suffix,fail_first", [(".pdf", False), (".html", True)])
def test_native_job_flow_keeps_snapshot_and_restores_controls(make_workspace, tmp_path, suffix, fail_first):
    entered = threading.Event()
    release = threading.Event()
    calls = []
    output = tmp_path / "output.xlsx"
    result = ExtractionResult(
        metadata={"pages_selected": "12,14-16"}, raw_text_rows=[], raw_table_cells=[], normalized_table_cells=[],
        extraction_audit_rows=[], statements={}, parsed_cells=[], unmapped_rows=[],
        financial_audit_rows=[{"Status": "WARN", "Check": "Example", "Scope": "Document", "Detail": "Review source"}],
    )

    def runner(source, pages, metadata, destination):
        calls.append((source, pages, metadata, destination))
        entered.set()
        assert release.wait(5), "Test did not release worker"
        if fail_first and len(calls) == 1:
            raise PermissionError("Output folder is not writable")
        return result, output

    app, dialogs = make_workspace(runner)
    source = tmp_path / f"filing{suffix}"
    source.write_bytes(b"Synthetic fixture for an injected runner")
    app.pdf_path.set(str(source))
    app.client_name.set("Original company")
    app.pdf_run_button.invoke()
    try:
        assert entered.wait(5)
        assert app.pdf_run_button.instate(["disabled"])
        assert app.lifecycle.phase is JobPhase.RUNNING
        assert ("Parsing HTML" if suffix == ".html" else "Scanning PDF") in app.pdf_status.get("1.0", "end")
        app.client_name.set("Next company")
        app._start_pdf_extraction()
        assert dialogs[-1] == ("info", "Busy", "An extraction is already running.")
    finally:
        release.set()
        app.controller.worker.join(5)
    assert not app.controller.is_running
    # The worker is done but its result has not reached the UI yet.
    assert app.lifecycle.phase is JobPhase.RUNNING
    app._start_pdf_extraction()
    assert len(calls) == 1
    app._poll_queue()

    assert calls[0][2]["client_name"] == "Original company"
    assert calls[0][1] == "auto"
    assert app.state.client_name == "Next company"
    assert app.pdf_run_button.instate(["!disabled"])
    assert app.shell.draft_notice.get() == "Changes not yet extracted"
    assert "modified" in app.shell.setup_panel.details_toggle.cget("text")
    if fail_first:
        assert app.lifecycle.phase is JobPhase.ERROR
        assert app.results.headline.get() == "Extraction failed"
        assert "Check access" in app.results.note.get("1.0", "end")
        assert app.pdf_run_button.cget("text") == "Retry"
        assert "ERROR: Output folder is not writable" in app.pdf_status.get("1.0", "end")
        app.pdf_run_button.invoke()
        app.controller.worker.join(5)
        assert not app.controller.is_running
        app._poll_queue()
        assert calls[1][2]["client_name"] == "Next company"
        assert app.shell.draft_notice.get() == ""
        assert "modified" not in app.shell.setup_panel.details_toggle.cget("text")

    assert app.pdf_run_button.cget("text") == "Open workbook"
    assert app.results.output.get() == str(output)
    assert app.lifecycle.phase is JobPhase.SUCCEEDED
    assert app.pdf_run_button.instate(["!disabled"])
    summary = app.pdf_status.get("1.0", "end")
    assert f"Saved workbook: {output}" in summary
    assert "Source pages: 12,14-16" in summary
    assert "Parsed statements:\n  None" in summary
    assert "[WARN] Example | Document: Review source" in summary

    next_source = tmp_path / "next.html"
    next_source.write_text("<html></html>", encoding="utf-8")
    app._set_pdf_path(next_source)
    assert app.lifecycle.phase is JobPhase.READY
    assert app.pdf_status.get("1.0", "end") == summary
    assert app.result_context.get() == f"Result from: {source}"
    assert app.shell.draft_notice.get() == "Changes not yet extracted"


def test_large_selection_review_cancel_and_html_keeps_pdf_draft(make_workspace, tmp_path, monkeypatch):
    from financial_statement_extract.page_selection import build_page_plan
    from models import PageSelectionReviewRequired

    def preflight(*args):
        raise PageSelectionReviewRequired(build_page_plan(args[0], [], list(range(1, 22)), "manual"))

    app, _dialogs = make_workspace(preflight)
    confirmations = []
    monkeypatch.setattr("app.messagebox.askyesnocancel", lambda *a, **k: confirmations.append(a) or None)
    source = tmp_path / "filing.pdf"
    source.write_bytes(b"Synthetic fixture for an injected runner")
    app.pdf_path.set(str(source))
    app._on_pages_focus_in()
    app.pages.set("1-21")
    app.pdf_run_button.invoke()
    app.controller.worker.join(5)
    app._poll_queue()
    assert len(confirmations) == 1
    assert not app.controller.is_running
    assert "no workbook created" in app.pdf_status.get("1.0", "end")
    assert app.lifecycle.phase is JobPhase.READY
    assert not app.lifecycle.awaiting_review
    assert app.pdf_run_button.instate(["!disabled"])
    assert app.shell.draft_notice.get() == "Changes not yet extracted"

    requests = []
    app.controller = SimpleNamespace(is_running=False, start=requests.append)
    html = tmp_path / "filing.html"
    html.write_text("<html></html>", encoding="utf-8")
    app.pdf_path.set(str(html))
    app.pdf_run_button.invoke()
    assert len(confirmations) == 1
    assert requests[0].pages == "1-21"
    assert app.state.pages == "1-21"


@pytest.mark.parametrize("choice, policy", [(True, "suggested"), (False, "exact")])
def test_page_review_resumes_only_with_explicit_policy(make_workspace, tmp_path, monkeypatch, choice, policy):
    from financial_statement_extract.page_selection import build_page_plan

    app, _dialogs = make_workspace(unexpected_runner)
    app.pdf_path.set(str(tmp_path / "filing.pdf"))
    request = app.state.to_request()
    app.lifecycle.begin(request, app.state.draft_key())
    plan = build_page_plan(request.input_path, [], list(range(1, 14)), "manual")
    requests = []
    app.controller = SimpleNamespace(is_running=False, start=requests.append)
    monkeypatch.setattr("app.messagebox.askyesnocancel", lambda *a, **k: choice)
    app._review_page_selection(plan, request)
    assert len(requests) == 1
    assert requests[0].page_policy == policy
    assert requests[0].pages == request.pages


def test_contextual_setup_and_disclosures_retain_values(make_workspace, tmp_path):
    app, _dialogs = make_workspace(unexpected_runner)
    pdf = tmp_path / "filing.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    html = tmp_path / "filing.html"
    html.write_text("<html></html>", encoding="utf-8")
    app.root.deiconify()
    app.root.update()
    app._set_pdf_path(pdf)
    app.pages.set("12,14-16")
    panel = app.shell.setup_panel
    panel.details_toggle.invoke()
    app.client_name.set("Example company")
    app.year.set("2026")
    app.period.set("Annual")
    app.audit_status.set("Company Prepared")
    app.root.update()
    client_control = panel.entries["client_name"]
    client_control.focus_force()
    panel.details_toggle.invoke()
    app.root.update()
    assert app.root.focus_get() is panel.details_toggle
    assert not panel.details_form.winfo_ismapped()
    assert panel.details_toggle.cget("text") == "Show details (modified)"
    app.shell.toggle_setup()
    app.root.update()
    assert not app.shell.setup.winfo_ismapped()
    assert app.shell.setup_toggle.cget("text") == "Show setup (modified)"
    assert app.shell.context.winfo_ismapped()
    assert app.shell.output_entry.winfo_ismapped()
    app.shell.toggle_setup()
    app._set_pdf_path(html)
    app.root.update()
    assert not panel.pages_field.winfo_ismapped()
    assert "does not apply to HTML" in panel.pages_help.cget("text")
    assert app.state.pages == "12,14-16"
    app._set_pdf_path(pdf)
    panel.details_toggle.invoke()
    app.root.update()
    assert panel.pages_field.winfo_ismapped()
    assert panel.entries["client_name"] is client_control
    assert app.state.metadata() == {"client_name": "Example company", "year": "2026",
                                    "period": "Annual", "audit_status": "Company Prepared"}
    assert app.state.pages == "12,14-16"
    assert app.controller.worker is None


def test_live_minimum_layout_keeps_context_actions_and_expanded_details_visible(make_workspace, tmp_path):
    app, _dialogs = make_workspace(unexpected_runner)
    source = tmp_path / ("Long company and filing description " * 4 + ".pdf")
    source.write_bytes(b"%PDF-1.4")
    app._set_pdf_path(source)
    app.root.geometry("640x700")
    app.root.deiconify()
    app.root.update()
    assert app.shell.narrow
    app.shell.toggle_setup()
    app.shell.setup_panel.details_toggle.invoke()
    app.root.update()
    assert app.shell.setup_visible
    for widget in (app.shell.primary_button, app.shell.status_label, app.shell.output_entry,
                   app.shell.source_entry, app.shell.setup_toggle, *app.shell.setup_panel.entries.values()):
        assert widget.winfo_ismapped(), str(widget)
        assert app.root.winfo_rootx() <= widget.winfo_rootx()
        assert app.root.winfo_rooty() <= widget.winfo_rooty()
        assert widget.winfo_rootx() + widget.winfo_width() <= app.root.winfo_rootx() + app.root.winfo_width()
        assert widget.winfo_rooty() + widget.winfo_height() <= app.root.winfo_rooty() + app.root.winfo_height()
    assert app.results.note.winfo_height() >= 60, "\n".join(
        str((str(w), w.winfo_width(), w.winfo_height(), w.winfo_reqheight()))
        for w in (app.shell, app.shell.context, app.shell.body, app.shell.setup,
                  app.shell.surface, app.shell.setup_panel, app.shell.setup_panel.details_form,
                  *app.shell.surface.winfo_children()))


def test_chooser_cancel_invalid_paths_and_missing_source_preserve_draft(make_workspace, tmp_path, monkeypatch):
    app, dialogs = make_workspace(unexpected_runner)
    source = tmp_path / "filing.pdf"
    source.write_bytes(b"%PDF-1.4")
    monkeypatch.setattr("app.filedialog.askopenfilename", lambda **_kw: str(source))
    app.pdf_run_button.invoke()
    assert app.state.input_path == str(source)
    assert app.pdf_run_button.cget("text") == "Extract to Excel"
    monkeypatch.setattr("app.filedialog.askdirectory", lambda **_kw: str(tmp_path / "destination"))
    app._choose_output_dir()
    draft = app.state.draft_key()
    monkeypatch.setattr("app.filedialog.askopenfilename", lambda **_kw: "")
    monkeypatch.setattr("app.filedialog.askdirectory", lambda **_kw: "")
    app._choose_pdf()
    app._choose_output_dir()
    assert app.state.draft_key() == draft
    assert not app._set_pdf_path(tmp_path / "missing.pdf")
    assert app.state.draft_key() == draft
    app.output_dir.set("//server/share/output")
    assert app.pdf_run_button.instate(["disabled"])
    assert app.shell.disabled_reason.get()
    app.output_dir.set(str(tmp_path))
    assert app.pdf_run_button.instate(["!disabled"])
    source.unlink()  # This test-created file disappears after it was selected.
    app.pdf_run_button.invoke()
    assert dialogs[-1][1] == "Unsupported file"
    assert not app._job_active
    assert app.controller.worker is None


def test_live_html_extraction_creates_workbook_and_clears_submitted_changes(make_workspace, tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "sample.html"
    source.write_text("""<html><body><h2>Consolidated Statements of Income</h2>
    <table><tr><th></th><th>2026</th><th>2025</th></tr>
    <tr><td>Revenue</td><td>300</td><td>250</td></tr>
    <tr><td>Cost of sales</td><td>180</td><td>150</td></tr>
    <tr><td>Gross profit</td><td>120</td><td>100</td></tr>
    <tr><td>Operating income</td><td>50</td><td>40</td></tr>
    <tr><td>Net income</td><td>30</td><td>20</td></tr></table></body></html>""", encoding="utf-8")
    app, dialogs = make_workspace(None)
    app._set_pdf_path(source)
    app.pages.set("12,14-16")  # Retained PDF draft must not constrain HTML.
    app.client_name.set("Synthetic example")
    app.pdf_run_button.invoke()
    app.controller.worker.join(20)
    assert not app.controller.is_running
    app._poll_queue()
    assert app.lifecycle.phase is JobPhase.SUCCEEDED, dialogs
    assert app.results.has_summary
    workbooks = list(tmp_path.glob("*.xlsx"))
    assert len(workbooks) == 1
    workbook = load_workbook(workbooks[0], read_only=True)
    try:
        values = [cell for sheet in workbook for row in sheet.values for cell in row]
        assert "Consolidated Statements of Income" in values
        assert "Revenue" in values
        assert 300 in values
    finally:
        workbook.close()
    assert "Consolidated Statements of Income" in app.pdf_status.get("1.0", "end")
    assert app.shell.draft_notice.get() == ""
    assert app.state.pages == "12,14-16"


def test_worker_start_failure_is_an_error_with_retryable_inputs(make_workspace, tmp_path, monkeypatch):
    app, dialogs = make_workspace(unexpected_runner)
    source = tmp_path / "sample.pdf"
    source.write_bytes(b"Synthetic fixture")
    app._set_pdf_path(source)
    app.client_name.set("Keep this draft")
    draft_key = app.state.draft_key()

    def fail_start(_request):
        raise RuntimeError("Worker unavailable")

    monkeypatch.setattr(app.controller, "start", fail_start)
    app.pdf_run_button.invoke()
    assert app.lifecycle.phase is JobPhase.ERROR
    assert app.lifecycle.error == "Worker unavailable"
    assert app.lifecycle.request.client_name == "Keep this draft"
    assert app.state.draft_key() == draft_key
    assert app.pdf_run_button.instate(["!disabled"])
    assert "Failed extraction" in app.result_context.get()
    assert "Worker unavailable" in app.pdf_status.get("1.0", "end")
    assert app.results.headline.get() == "Extraction failed"
    assert app.pdf_run_button.cget("text") == "Retry"
    app._refresh_setup()
    assert app.lifecycle.phase is JobPhase.ERROR
    app.client_name.set("Corrected draft")
    assert app.lifecycle.phase is JobPhase.READY
    assert app.lifecycle.request.client_name == "Keep this draft"


def test_review_resume_failure_restores_action_and_retains_request(make_workspace, tmp_path, monkeypatch):
    from financial_statement_extract.page_selection import build_page_plan

    app, dialogs = make_workspace(unexpected_runner)
    app.pdf_path.set(str(tmp_path / "sample.pdf"))
    app.pages.set("1-30")
    request = app.state.to_request()
    app.lifecycle.begin(request, app.state.draft_key())
    plan = build_page_plan(request.input_path, [], list(range(1, 31)), "manual")

    def review(*_args, **_kwargs):
        assert app.lifecycle.phase is JobPhase.RUNNING
        assert app.lifecycle.awaiting_review
        assert app.pdf_run_button.instate(["disabled"])
        assert "Review" in app.shell.disabled_reason.get()
        app.client_name.set("Next request")
        return False

    def fail_resume(_request):
        raise RuntimeError("Could not resume worker")

    monkeypatch.setattr("app.messagebox.askyesnocancel", review)
    monkeypatch.setattr(app.controller, "start", fail_resume)
    app._review_page_selection(plan, request)
    assert app.lifecycle.phase is JobPhase.ERROR
    assert not app.lifecycle.awaiting_review
    assert app.lifecycle.request.page_policy == "exact"
    assert app.lifecycle.request.client_name == ""
    assert app.state.client_name == "Next request"
    assert app.pdf_run_button.instate(["!disabled"])
    assert app.results.headline.get() == "Extraction failed"
    assert "Could not resume worker" in app.results.diagnostics.get("1.0", "end")


def test_fake_runner_exercises_rapid_retries_without_loading_extraction(tmp_path):
    import subprocess
    import sys

    # A fresh process proves these lifecycle checks never rely on a loaded
    # pipeline, document parser, or a real workbook. Paths are synthetic only.
    completed = subprocess.run([sys.executable, "-c", """
import sys
from financial_statement_extract.ui.controller import ExtractionController
from financial_statement_extract.ui.lifecycle import JobLifecycle, JobPhase
from financial_statement_extract.ui.state import WorkspaceState
calls = []
def fake_runner(*args):
    calls.append(args)
    if len(calls) % 2:
        raise PermissionError('Simulated output failure')
    return object(), 'not-created.xlsx'
controller = ExtractionController(fake_runner)
draft = WorkspaceState(input_path='synthetic.pdf')
job = JobLifecycle()
for index in range(20):
    job.sync_draft(draft)
    job.begin(draft.to_request(), draft.draft_key())
    controller.start(job.request)
    controller.worker.join(5)
    assert not controller.is_running
    assert job.phase is JobPhase.RUNNING
    event = controller.work_queue.get_nowait()
    if event[0] == 'pdf_error':
        job.fail(event[1])
        assert job.phase is JobPhase.ERROR
    else:
        job.succeed()
        assert job.phase is JobPhase.SUCCEEDED
assert not {'pipeline', 'pandas', 'pdfplumber', 'camelot'} & sys.modules.keys()
assert len(calls) == 20
"""], capture_output=True, text=True, timeout=20)
    assert completed.returncode == 0, completed.stderr


@pytest.mark.parametrize("fail", [False, True])
def test_progress_keeps_ui_responsive_and_is_rendered_only_on_ui_thread(make_workspace, tmp_path, monkeypatch, fail):
    from dataclasses import replace
    from financial_statement_extract.progress import ExtractionStage, ProgressEvent

    entered = threading.Event()
    release = threading.Event()
    stage = ProgressEvent(ExtractionStage.WRITING, "Synthetic workbook stage")
    result = ExtractionResult({}, [], [], [], [], {}, [], [], [])

    def runner(*_args, progress):
        progress(stage)
        entered.set()
        assert release.wait(5)
        if fail:
            raise PermissionError("Synthetic write failure")
        return result, tmp_path / "not-created.xlsx"

    app, dialogs = make_workspace(runner, reports_progress=True)
    source = tmp_path / "synthetic.pdf"
    source.write_bytes(b"Synthetic placeholder, not parsed")
    app._set_pdf_path(source)
    main_thread = threading.get_ident()
    original_status = app.shell.set_status
    seen = []

    def status(*args):
        assert threading.get_ident() == main_thread
        seen.append(args[0])
        original_status(*args)

    monkeypatch.setattr(app.shell, "set_status", status)
    app.pdf_run_button.invoke()
    try:
        assert entered.wait(5)
        app._poll_queue()
        assert app.lifecycle.progress is stage
        assert "Writing workbook" in app.shell.status_text.get()
        assert str(app.progress.cget("mode")) == "indeterminate"
        app.client_name.set("Next run")
        tick = []
        app.root.after_idle(lambda: tick.append(True))
        app.root.update()
        assert tick == [True]
        assert app.controller.is_running
        assert app.pdf_run_button.instate(["disabled"])
        monkeypatch.setattr("app.messagebox.askyesno", lambda *_args, **_kw: False)
        app._on_close()
        assert not app.closing and app.controller.is_running
        # A stale report from a different request cannot replace this stage.
        stale = ProgressEvent(ExtractionStage.PARSING, "Stale event")
        app.controller.work_queue.put(("progress", stale, replace(app.lifecycle.request)))
        app._poll_queue()
        assert app.lifecycle.progress is stage
    finally:
        release.set()
        app.controller.worker.join(5)
    app._poll_queue()
    assert app.lifecycle.phase is (JobPhase.ERROR if fail else JobPhase.SUCCEEDED)
    assert app.results.headline.get() == ("Extraction failed" if fail else "Workbook saved · review required")
    assert app.state.client_name == "Next run"
    terminal_status = app.shell.status_text.get()
    app.controller.work_queue.put(("progress", stage, app.lifecycle.request))
    app._poll_queue()
    assert app.shell.status_text.get() == terminal_status
    assert any("Writing workbook" in text for text in seen)


def test_result_actions_use_saved_location_and_recover_from_missing_workbook(make_workspace, tmp_path, monkeypatch):
    output = tmp_path / "saved.xlsx"
    output.touch()
    result = ExtractionResult({}, [], [], [], [], {}, [], [], [])
    app, dialogs = make_workspace(lambda *_args: (result, output))
    source = tmp_path / "source.html"
    source.write_text("Synthetic runner placeholder", encoding="utf-8")
    app._set_pdf_path(source)
    app.pdf_run_button.invoke()
    app.controller.worker.join(5)
    app._poll_queue()
    assert not dialogs  # Success is an inline surface, not a blocking popup.
    assert app.pdf_run_button.cget("text") == "Open workbook"
    opened = []
    monkeypatch.setattr("financial_statement_extract.ui.file_actions.os.startfile", lambda *args: opened.append(args))
    app.pdf_run_button.invoke()
    assert opened == [(str(output), "open")]
    app.output_dir.set(str(tmp_path / "next-output"))
    assert app.pdf_run_button.cget("text") == "Extract to Excel"
    assert app.result_action.cget("text") == "Open workbook"
    app.folder_button.invoke()
    assert opened[-1] == (str(tmp_path), "open")
    output.unlink()  # Test-owned placeholder only.
    app.result_action.invoke()
    assert len(opened) == 2
    assert dialogs[-1][1] == "Could not open result"
    assert app.results.has_summary
    assert app.results.output.get() == str(output)
    assert app.state.output_dir == str(tmp_path / "next-output")
    assert "Workbook is no longer available" in app.pdf_status.get("1.0", "end")


def test_results_tabs_diagnostics_and_long_checks_retain_selection_at_minimum_size(make_workspace, tmp_path):
    result = ExtractionResult({"pages_selected": "12,14-16"}, [], [], [], [], {}, [], [], [])
    result.financial_audit_rows = [dict(Status="WARN", Check=f"Long check {i}", Scope="Document",
                                      Detail="Long source-preserving diagnostic " * 25) for i in range(45)]
    app, _dialogs = make_workspace(lambda *_args: (result, tmp_path / "saved.xlsx"))
    source = tmp_path / "source.pdf"
    source.write_bytes(b"Synthetic runner placeholder")
    app._set_pdf_path(source)
    app.pdf_run_button.invoke()
    app.controller.worker.join(5)
    app._poll_queue()
    app.root.geometry("640x700")
    app.root.deiconify()
    app.root.update()
    view = app.results
    tree = view.tables["Checks"]
    assert view.notebook.index(view.notebook.select()) == 1
    tree.selection_set("44")
    tree.focus("44")
    tree.see("44")
    app.root.update()
    scroll = tree.yview()
    assert "Long check 44" in view.details["Checks"].get("1.0", "end")
    assert str(view.diagnostics.cget("state")) == "disabled"
    tree.focus_force()
    view.toggle.invoke()
    app.root.update()
    assert app.root.focus_get() is view.toggle
    assert view.diagnostics_visible
    view.toggle.invoke()
    app.root.update()
    assert tree.selection() == ("44",)
    assert tree.yview() == scroll
    for widget in (view.toggle, app.folder_button, app.pdf_run_button, app.shell.status_label, tree):
        assert widget.winfo_ismapped()
        assert widget.winfo_rootx() + widget.winfo_width() <= app.root.winfo_rootx() + app.root.winfo_width()
        assert widget.winfo_rooty() + widget.winfo_height() <= app.root.winfo_rooty() + app.root.winfo_height()
    assert tree.winfo_height() >= 70
