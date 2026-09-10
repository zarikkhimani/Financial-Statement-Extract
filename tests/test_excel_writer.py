from openpyxl import load_workbook
import pandas as pd

from excel_writer import write_extraction_workbook
from models import ExtractionResult
from pipeline import parse_text_to_result


def test_pasted_text_result_writes_workbook(tmp_path):
    result = parse_text_to_result(
        """Example Company
Income Statements
Years Ended December 31, 2026 2025
Revenue 300 250
Net income 30 20
"""
    )

    output = write_extraction_workbook(result, tmp_path / "result.xlsx")

    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["Income Statement"]
        sheet = workbook["Income Statement"]
        assert sheet["A1"].value == "Example Company"
        assert sheet["A2"].value == "Income Statements"
        assert sheet["A3"].value is None
        assert sheet["B5"].value == "2026"
        assert sheet["A6"].value == "Revenue"
        assert sheet.freeze_panes is None
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.fill is None or cell.fill.fill_type is None
                if cell.font is not None and cell.font.color is not None and cell.font.color.type == "rgb":
                    assert cell.font.color.rgb in {"00000000", "FF000000"}
    finally:
        workbook.close()

def test_reported_decimal_precision_is_preserved(tmp_path):
    result = parse_text_to_result(
        """Example Company
Income Statements
Years Ended December 31, 2026 2025
Revenue 300.5 250.0
Net income 30.2 20.1
"""
    )

    output = write_extraction_workbook(result, tmp_path / "decimal_result.xlsx")

    workbook = load_workbook(output, read_only=True)
    try:
        sheet = workbook["Income Statement"]
        assert sheet["B6"].value == 300.5
        assert sheet["B6"].number_format == "#,##0.0;(#,##0.0);-"
    finally:
        workbook.close()

def test_period_columns_are_wide_enough_for_full_dates(tmp_path):
    result = parse_text_to_result(
        """Example Fund
Consolidated Statements of Assets and Liabilities
December 31, 2025 December 31, 2024
Assets
Cash and cash equivalents 1234567890 100000
Total Assets 1234567890 100000
"""
    )

    output = write_extraction_workbook(result, tmp_path / "dated_result.xlsx")

    workbook = load_workbook(output)
    try:
        sheet = workbook["Balance Sheet"]
        period_columns = next(
            dimension
            for dimension in sheet.column_dimensions.values()
            if dimension.min <= 2 and dimension.max >= 3
        )
        assert period_columns.width >= 19
    finally:
        workbook.close()


def test_long_interim_period_headers_wrap_and_have_sufficient_height(tmp_path):
    result = parse_text_to_result(
        """Example Bank
Consolidated Statements of Comprehensive Income (unaudited)
Three Months Ended June 30, Six Months Ended June 30,
(in millions) 2025 2024 2025 2024
Net income 15 14 30 27
"""
    )

    output = write_extraction_workbook(result, tmp_path / "interim_result.xlsx")
    workbook = load_workbook(output)
    try:
        sheet = workbook["Income Statement"]
        assert sheet["B5"].alignment.wrap_text is True
        assert sheet.row_dimensions[5].height >= 30
    finally:
        workbook.close()


def test_pdf_result_writes_all_three_raw_methods_as_editable_grids(tmp_path):
    result = parse_text_to_result(
        """Example Company
Income Statements
Years Ended December 31, 2026 2025
Revenue 300 250
Net income 30 20
"""
    )
    result.metadata["source_file"] = "example.pdf"
    result.raw_table_cells = [
        {
            "source_page": 2,
            "table_id": "P0002_LATTICE_01",
            "flavor": "lattice",
            "selected": True,
            "table_score": 95.0,
            "accuracy": 99.0,
            "whitespace": 1.0,
            "row_index": 0,
            "column_index": 0,
            "raw_text": "Revenue",
        },
        {
            "source_page": 2,
            "table_id": "P0002_LATTICE_01",
            "flavor": "lattice",
            "selected": True,
            "table_score": 95.0,
            "accuracy": 99.0,
            "whitespace": 1.0,
            "row_index": 0,
            "column_index": 1,
            "raw_text": "=1+1",
        },
        {
            "source_page": 2,
            "table_id": "P0002_STREAM_01",
            "flavor": "stream",
            "selected": False,
            "row_index": 0,
            "column_index": 0,
            "raw_text": "Net income",
        },
        {
            "source_page": 2,
            "table_id": "P0002_PDFPLUMBER_01",
            "flavor": "pdfplumber",
            "selected": False,
            "row_index": 0,
            "column_index": 0,
            "raw_text": "Cash",
        },
    ]

    output = write_extraction_workbook(result, tmp_path / "raw_methods.xlsx")
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == [
            "Income Statement",
            "Raw Camelot Lattice",
            "Raw Camelot Stream",
            "Raw PDFPlumber",
        ]
        lattice = workbook["Raw Camelot Lattice"]
        assert lattice["A4"].value == "Table: P0002_LATTICE_01"
        assert lattice["A5"].value == "Revenue"
        assert lattice["B5"].value == "=1+1"
        assert lattice["B5"].data_type == "s"
        assert workbook["Raw Camelot Stream"]["A5"].value == "Net income"
        assert workbook["Raw PDFPlumber"]["A5"].value == "Cash"
        for sheet_name in workbook.sheetnames[1:]:
            sheet = workbook[sheet_name]
            assert sheet.freeze_panes is None
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.fill is None or cell.fill.fill_type is None
                    if cell.font is not None and cell.font.color is not None and cell.font.color.type == "rgb":
                        assert cell.font.color.rgb in {"00000000", "FF000000"}
    finally:
        workbook.close()


def test_document_text_is_always_written_as_literal_strings(tmp_path):
    first_period = "@SOURCE_PERIOD"
    second_period = "https://example.invalid/period"
    statement = pd.DataFrame(
        [
            {
                "RowType": "Section",
                "RawItem": "=SECTION()",
                first_period: None,
                second_period: None,
            },
            {
                "RowType": "Data",
                "RawItem": '=HYPERLINK("https://example.invalid", "open")',
                first_period: 125.5,
                second_period: 100,
            },
            {
                "RowType": "Data",
                "RawItem": "https://example.invalid/line-item",
                first_period: 1,
                second_period: 2,
            },
        ]
    )
    statement.attrs.update(
        {
            "company_name": "=1+1",
            "statement_title": "+SUM(1,1)",
            "raw_unit_note": "-1+1",
            "period_labels": [first_period, second_period],
        }
    )
    result = ExtractionResult(
        metadata={"source_file": "example.html"},
        raw_text_rows=[],
        raw_table_cells=[],
        normalized_table_cells=[],
        extraction_audit_rows=[],
        statements={"IncomeStatement": statement},
        parsed_cells=[],
        financial_audit_rows=[],
        unmapped_rows=[],
    )

    output = write_extraction_workbook(result, tmp_path / "literal_strings.xlsx")
    workbook = load_workbook(output, data_only=False)
    try:
        sheet = workbook["Income Statement"]
        expected_strings = {
            "A1": "=1+1",
            "A2": "+SUM(1,1)",
            "A3": "-1+1",
            "B5": first_period,
            "C5": second_period,
            "A6": "=SECTION()",
            "A7": '=HYPERLINK("https://example.invalid", "open")',
            "A8": "https://example.invalid/line-item",
        }
        for coordinate, expected in expected_strings.items():
            cell = sheet[coordinate]
            assert cell.value == expected
            assert cell.data_type == "s"
            assert cell.hyperlink is None
        assert sheet["B7"].value == 125.5
        assert sheet["B7"].data_type == "n"
    finally:
        workbook.close()
