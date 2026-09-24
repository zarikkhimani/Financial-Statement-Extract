from openpyxl import load_workbook
import pandas as pd

from audit import _period_columns, _series_value
from excel_writer import write_extraction_workbook
from mapper import map_concept
from pipeline import parse_text_to_result
from structure import detect_schedule_value_headers


def test_investment_company_result_is_related_not_a_corporate_synonym():
    source_label = "Net increase in net assets from operations"
    fund = map_concept("IncomeStatement", source_label)
    corporate = map_concept("IncomeStatement", "Net income")
    assert fund.raw_item == source_label
    assert fund.internal_id != corporate.internal_id
    assert fund.standard_item != corporate.standard_item
    assert fund.analytical_family == corporate.analytical_family == "period_result"
    assert fund.relationship == corporate.relationship == "exact_concept"


def test_related_and_inferred_mappings_cannot_satisfy_equality_checks():
    frame = pd.DataFrame([
        {"StandardItem": "Net Income (Loss)", "MappingRelationship": "related", "2026": 30,
         "InternalID": "corporate_net_income", "AnalyticalFamily": "period_result"},
        {"StandardItem": "Total Assets", "MappingRelationship": "inferred", "2026": 500},
    ])
    assert _period_columns(frame) == ["2026"]
    assert _series_value(frame, "Net Income (Loss)", "2026") is None
    assert _series_value(frame, "Total Assets", "2026") is None


def test_source_notes_punctuation_and_labels_reach_excel(tmp_path):
    result = parse_text_to_result("""Example Fund
Consolidated Statements of Operations (unaudited)
2026 2025
Other income (expense):
Interest income (3) 10 9
Net increase in net assets from operations 30 20
""")
    frame = result.statements["IncomeStatement"]
    assert "Interest income (3)" in set(frame.RawItem)
    row = frame[frame.RawItem == "Net increase in net assets from operations"].iloc[0]
    assert row.InternalID == "investment_company_net_increase_from_operations"
    output = write_extraction_workbook(result, tmp_path / "fund.xlsx")
    book = load_workbook(output)
    try:
        sheet = book["Statement 1"]
        assert sheet["C5"].value == "Consolidated Statements of Operations (unaudited)"
        labels = [cell.value for cell in sheet["C"]]
        assert "Other income (expense):" in labels
        assert "Interest income (3)" in labels
        assert "Net increase in net assets from operations" in labels
        assert row.InternalID not in labels
        assert row.StandardItem not in labels
    finally:
        book.close()


def test_unnamed_totals_keep_values_without_invented_display_labels(tmp_path):
    result = parse_text_to_result("""Example Company
Balance Sheets
2026 2025
Current assets
Cash 60 50
Accounts receivable 40 30
100 80
""")
    frame = result.statements["BalanceSheet"]
    inferred = frame[frame.MappingRelationship == "inferred"].iloc[0]
    assert inferred.RawItem == ""
    assert inferred["2026"] == 100
    output = write_extraction_workbook(result, tmp_path / "blank-label.xlsx")
    book = load_workbook(output)
    try:
        sheet = book["Balance Sheets"]
        assert any(row[0].value is None and row[1].value == 100 for row in sheet.iter_rows(min_col=3))
        assert "Total current assets" not in [cell.value for cell in sheet["C"]]
    finally:
        book.close()


def test_schedule_headers_keep_source_case_wording_and_order():
    assert detect_schedule_value_headers([
        {"raw_text": "Investment FAIR VALUE cost Percent of Net Assets"},
    ]) == ["FAIR VALUE", "cost", "Percent of Net Assets"]
