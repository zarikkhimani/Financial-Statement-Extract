import json

from openpyxl import Workbook
import pytest

from scripts.compare_workbook_baseline import (
    amount, compare_rows, file_hash, has_differences, inspect_workbook, read_statement, verify_baseline,
)


def row(label, *values):
    return {"label": label, "values": list(values)}


def test_legitimate_equal_rows_are_counted_before_excess_duplicates():
    investment = row("Same investment", 100, 99)
    expected = [investment, investment]
    result = compare_rows(expected, [investment, investment, investment])
    assert result["matched_rows"] == 2
    assert len(result["duplicated_rows"]) == 1
    assert not result["missing_rows"]


def test_duplicate_cannot_conceal_a_missing_row_even_with_equal_totals():
    a, b = row("A", 100), row("B", 100)
    result = compare_rows([a, b], [a, a])
    assert result["missing_rows"] == [b]
    assert result["duplicated_rows"] == [a]


def test_changed_amounts_are_detected_even_when_sum_is_unchanged():
    result = compare_rows([row("Revenue", 100, 200)], [row("Revenue", 110, 190)])
    assert result["changed_rows"][0]["changed_fields"] == [0, 1]
    assert not result["missing_rows"]


def test_missing_amount_is_not_zero_but_em_dash_and_blank_are_missing():
    result = compare_rows([row("A", None, 2)], [row("A", "\u2014", 2)])
    assert result["matched_rows"] == 1
    changed = compare_rows([row("A", None, 2)], [row("A", 0, 2)])
    assert changed["changed_rows"][0]["changed_fields"] == [0]


@pytest.mark.parametrize("value", [True, "NaN", "Infinity", "not an amount"])
def test_invalid_amounts_fail_clearly(value):
    with pytest.raises(ValueError):
        amount(value)


def test_incomplete_numeric_columns_are_not_reported_as_changed_numbers():
    result = compare_rows([row("Balance", 100, 200, 300)], [row("Balance", 300)])
    assert len(result["unstructured_rows"]) == 1
    assert not result["changed_rows"]


def test_adjacent_long_label_does_not_steal_values_from_short_label():
    book = Workbook()
    sheet = book.active
    sheet.append([None])
    sheet["A5"], sheet["B5"], sheet["C5"] = "Revenue", 10, 20
    sheet["A6"], sheet["B6"], sheet["C6"] = "A much longer expense description", 3, 4
    ref = dict(id="income", fields=["2026", "2025"], headings=[], rows=[
        row("Revenue", 10, 20), row("A much longer expense description", 3, 4)])
    actual, _ = read_statement(sheet, ref)
    assert compare_rows(ref["rows"], actual)["matched_rows"] == 2


def test_year_header_does_not_become_cash_and_column_shifts_are_reported():
    book = Workbook()
    sheet = book.active
    sheet["C5"], sheet["F5"] = 2026, 2025
    sheet["A6"], sheet["D6"], sheet["G6"] = "Cash", 10, 20
    sheet["A7"], sheet["C7"], sheet["F7"] = "Other assets", 3, 4
    ref = dict(id="balance", fields=["2026", "2025"], headings=["2026", "2025"], rows=[
        row("Cash", 10, 20), row("Other assets", 3, 4)])
    actual, layout = read_statement(sheet, ref)
    assert len(actual) == 2
    assert compare_rows(ref["rows"], actual)["matched_rows"] == 2
    assert layout["inconsistent_columns"]
    assert not layout["missing_headings"]


def test_amounts_embedded_in_description_are_retained_as_unstructured():
    sheet = Workbook().active
    sheet["A5"] = "Loan A One stop 12/2030 (2)"
    ref = dict(id="schedule", fields=["Principal", "Cost", "Percent", "Value"], headings=[],
               rows=[row("Loan A", None, -2, None, None)])
    actual, _ = read_statement(sheet, ref)
    result = compare_rows(ref["rows"], actual)
    assert not result["missing_rows"]
    assert len(result["unstructured_rows"]) == 1


def test_missing_statement_is_explicit_and_read_does_not_modify_file(tmp_path):
    path = tmp_path / "source.xlsx"
    book = Workbook()
    book.save(path)
    before = file_hash(path)
    ref = dict(statement_sections=[dict(id="balance", title="Balance Sheets", pages=[1],
                                        fields=["2026"], headings=["2026"], numeric_count=1,
                                        rows=[row("Cash", 10)])])
    result = inspect_workbook(path, ref)
    assert result["statements"][0]["error"] == "Statement missing"
    assert file_hash(path) == before


def test_reference_integrity_failure_stops_comparison(tmp_path):
    source = tmp_path / "source.pdf"
    source.write_bytes(b"reference")
    expected = tmp_path / "expected.json"
    expected.write_text(json.dumps(dict(source_sha256=file_hash(source))), encoding="utf-8")
    manifest = {"files": {
        "source": {"file": "source.pdf", "sha256": file_hash(source)},
        "expected.json": {"file": "expected.json", "sha256": file_hash(expected)},
    }}
    (tmp_path / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    verify_baseline(tmp_path)
    source.write_bytes(b"modified")
    with pytest.raises(ValueError, match="integrity"):
        verify_baseline(tmp_path)


def test_reordered_rows_are_separate_from_changed_values():
    a, b = row("A", 10), row("B", 20)
    result = compare_rows([a, b], [b, a])
    assert result["matched_rows"] == 2
    assert result["order_changed"]
    assert not result["changed_rows"]


def test_replacement_glyph_is_distinct_from_normal_dash():
    sheet = Workbook().active
    sheet["A5"], sheet["B5"], sheet["C5"] = "A", "\u2014", 2
    sheet["A6"] = "\ufffd"
    ref = dict(id="income", fields=["2026", "2025"], headings=[], rows=[row("A", None, 2)])
    _, layout = read_statement(sheet, ref)
    assert layout["replacement_characters"] == ["A6"]


def test_total_formatting_difference_is_not_reported_as_clean():
    assert has_differences({"statements": [{"layout": {"totals_without_bold": [8]}}]})
