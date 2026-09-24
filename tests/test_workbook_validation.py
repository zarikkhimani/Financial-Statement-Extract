from copy import deepcopy
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook, Workbook

from audit import _series_value, audit_cross_statement, audit_units
from excel_writer import write_extraction_workbook
from pipeline import parse_text_to_result
from financial_statement_extract.workbook_validation import (
    audit_balance_grid, audit_cash_grids, audit_equity_grid, audit_recovered_statement,
    audit_selected_page_coverage, audit_source_cross_statement, audit_text_row_coverage, validate_grid,
    recovered_tables,
)
from scripts.compare_workbook_baseline import read_statement


def grid(rows, periods=('June 30, 2026', 'September 30, 2025'), flavor='balance_geometry'):
    return dict(table_id='test', flavor=flavor, rows=rows, row_kinds=['data'] * len(rows),
                row_sources=[[{'source_page': 1, 'line_no': i+1}] for i in range(len(rows))],
                period_labels=list(periods), source_page=1)


def balance():
    return grid([['Total Assets', '100', '90'], ['Total Liabilities', '60', '50'],
                 ['Total Net Assets', '40', '40'], ['Total Liabilities and Total Net Assets', '100', '90']])


def frame(table):
    f = pd.DataFrame()
    f.attrs.update(source_tables=[table], unit_label='thousands', currency='$')
    return f


def test_fund_equation_checks_both_component_sum_and_reported_total():
    table = balance()
    assert all(r['Status'] == 'PASS' for r in audit_balance_grid(table))
    table['rows'][2][1] = '42'
    checks = audit_balance_grid(table)
    assert checks[0]['Status'] == 'FAIL'  # matching printed totals cannot hide a wrong component
    assert checks[1]['Status'] == 'PASS'


@pytest.mark.parametrize('mutation', ['missing', 'duplicate', 'blank'])
def test_missing_ambiguous_or_blank_totals_cannot_pass(mutation):
    table = balance()
    if mutation == 'missing':
        table['rows'].pop(2)
        table['row_kinds'].pop(2)
    elif mutation == 'duplicate':
        table['rows'].append(table['rows'][2])
        table['row_kinds'].append('data')
    else:
        table['rows'][2][1] = ''
    assert audit_balance_grid(table)[0]['Status'] == 'NOT_TESTED'


@pytest.mark.parametrize('mutation', ['headers', 'width', 'metadata', 'invalid', 'nonfinite', 'source_duplicate', 'empty_source'])
def test_bad_column_or_source_structure_is_explicit_and_stops_financial_passes(mutation):
    table = balance()
    if mutation == 'headers':
        table['period_labels'][1] = table['period_labels'][0]
    elif mutation == 'width':
        table['rows'][0].pop()
    elif mutation == 'metadata':
        table['row_sources'].pop()
    elif mutation == 'invalid':
        table['rows'][0][1] = '10 90'
    elif mutation == 'nonfinite':
        table['rows'][0][1] = 'NaN'
    elif mutation == 'empty_source':
        table['row_sources'][0] = []
    else:
        table['row_sources'][1] = table['row_sources'][0]
    checks = audit_recovered_statement('BalanceSheet', [table])
    assert checks[0]['Status'] == 'FAIL'
    assert not any(r['Status'] == 'PASS' for r in checks)


def test_equal_values_at_distinct_source_positions_are_not_duplicate_rows():
    assert validate_grid(grid([['Interest', '10', '9'], ['Interest', '10', '9']]))['Status'] == 'PASS'


def test_an_untyped_source_grid_does_not_claim_to_be_a_verified_recovery():
    f = frame({'rows': [['Income', '10']]})
    assert recovered_tables(f, 'IncomeStatement') == []


def cash():
    return grid([['Cash, beginning of period', '100'], ['Net change in cash', '10'],
                 ['Effect of foreign currency exchange rates', '(2)'], ['Cash, end of period', '108']],
                periods=('Nine months ended June 30, 2026',), flavor='cash_flow_text')


def test_cash_rollforward_includes_fx_and_detects_changed_end():
    table = cash()
    assert audit_cash_grids([table])[0]['Status'] == 'PASS'
    table['rows'][-1][1] = '110'
    assert audit_cash_grids([table])[0]['Status'] == 'FAIL'


def test_cross_statement_never_pairs_same_year_with_different_dates_or_durations():
    checks = audit_source_cross_statement({'CashFlowStatement': frame(cash()), 'BalanceSheet': frame(balance())})
    assert not any(r['Status'] == 'PASS' for r in checks)
    assert checks[0]['Status'] == 'NOT_TESTED'
    a = pd.DataFrame([dict(StandardItem='Net Income (Loss)', MappingRelationship='exact_concept', **{'Three months ended June 30, 2026': 50})])
    b = pd.DataFrame([dict(StandardItem='Net Income (Loss)', MappingRelationship='exact_concept', **{'Nine months ended June 30, 2026': 50})])
    assert audit_cross_statement({'IncomeStatement': a, 'CashFlowStatement': b}) == []


def test_cross_statement_mismatched_units_are_not_tested():
    reconciliation = grid([['Cash', '10'], ['Total cash', '10']], periods=('June 30, 2026',), flavor='cash_flow_text')
    bs = frame(grid([['Cash', '10']], periods=('June 30, 2026',)))
    bs.attrs['unit_label'] = 'millions'
    checks = audit_source_cross_statement({'CashFlowStatement': frame(reconciliation), 'BalanceSheet': bs})
    assert checks[0]['Status'] == 'NOT_TESTED'


def test_equity_changes_and_movement_sums_are_both_checked():
    table = grid([['Balance at September 30, 2025', '10'], ['Issued', '2'], ['Redeemed', '—'],
                  ['Total increase for the nine months ended June 30, 2026', '2'], ['Balance at June 30, 2026', '12']],
                 periods=('Shares',), flavor='equity_geometry')
    table['component_labels'] = table.pop('period_labels')
    table['period_sections'] = [dict(start=0, end=5, caption='Nine months ended June 30, 2026')]
    assert all(r['Status'] == 'PASS' for r in audit_equity_grid(table))
    table['rows'][1][1] = '4'
    checks = audit_equity_grid(table)
    assert checks[0]['Status'] == 'PASS' and checks[1]['Status'] == 'FAIL'


def test_duplicate_exact_mapping_is_not_arbitrarily_reduced_to_first_row():
    f = pd.DataFrame([dict(StandardItem='Total Assets', MappingRelationship='exact_concept', **{'2026': 100})] * 2)
    assert _series_value(f, 'Total Assets', '2026') is None


def test_unknown_units_cannot_receive_a_pass():
    assert audit_units({'BalanceSheet': pd.DataFrame([{'2026': 10}])})[0]['Status'] == 'NOT_TESTED'


def test_selected_page_without_exported_table_is_a_failure():
    page = SimpleNamespace(page=3, statement_type='BalanceSheet', statement_types=())
    result = SimpleNamespace(page_plan=SimpleNamespace(pages=[page], selected_pages=[3]), statements={})
    assert audit_selected_page_coverage(result)[0]['Status'] == 'FAIL'


def test_source_row_coverage_catches_equal_total_missing_row_and_duplicate():
    table = grid([['Cash', '10', '9'], ['Other assets', '10', '9']])
    f = frame(table)
    result = SimpleNamespace(raw_text_rows=[dict(source_page=1, raw_text='Cash 10 9'),
                                            dict(source_page=1, raw_text='Other assets 10 9')])
    assert audit_text_row_coverage(result, f, 'BalanceSheet', 1)['Status'] == 'PASS'
    table['rows'][1] = deepcopy(table['rows'][0])
    check = audit_text_row_coverage(result, f, 'BalanceSheet', 1)
    assert check['Status'] == 'WARN' and 'other assets' in check['Detail']


def test_review_exposes_all_statuses_safe_literal_text_and_no_freeze(tmp_path):
    result = parse_text_to_result('Income Statements\n2026 2025\nRevenue 10 9')
    result.financial_audit_rows = [dict(Status=s, Check='=unsafe', Scope='Period', Detail='Details')
                                   for s in ['FAIL', 'NOT_TESTED', 'PASS', 'WARN']]
    result.extraction_audit_rows = []
    book = load_workbook(write_extraction_workbook(result, tmp_path/'review.xlsx'))
    sheet = book['Review']
    assert [sheet.cell(r, 3).value for r in range(9, 13)] == ['FAIL', 'WARN', 'NOT_TESTED', 'PASS']
    assert sheet['D9'].value == '=unsafe' and sheet['D9'].data_type == 's'
    assert '1 failed' in sheet['C5'].value and '1 not tested' in sheet['C5'].value
    assert all(w.freeze_panes is None for w in book)
    book.close()


def test_prose_footnote_is_not_an_extra_financial_row_but_note_with_values_is_checked():
    sheet = Workbook().active
    sheet.append([])
    sheet['A5'] = '(1) See Note 2 for Cash.'
    sheet.append(['(1) Cash', 10, 9])
    reference = dict(id='cash', fields=['2026', '2025'], headings=[], rows=[dict(label='Cash', values=[10, 9])])
    actual, layout = read_statement(sheet, reference)
    assert len(actual) == 1 and actual[0]['row'] == 6
    assert not layout['labels_with_extra_text']


def test_numbered_financial_row_with_embedded_values_is_not_hidden_as_prose():
    sheet = Workbook().active
    sheet['A5'] = '(1) Cash 10 9'
    reference = dict(id='cash', fields=['2026', '2025'], headings=[], rows=[dict(label='Cash', values=[10, 9])])
    actual, layout = read_statement(sheet, reference)
    assert len(actual) == 1 and actual[0]['values'] == []
    assert layout['labels_with_extra_text'] == [5]
