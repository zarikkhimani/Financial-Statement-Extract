from copy import deepcopy
import hashlib
import json

from openpyxl import Workbook
import pytest

from financial_statement_extract.workbook_validation import audit_balance_grid
from scripts.compare_apple_balance_baseline import load_reference
from scripts.validate_apple_balance import audit_saved_cells


def corporate():
    rows = [
        ['Current assets:', '', ''], ['Cash', '60', '50'], ['Other assets', '40', '30'],
        ['Total current assets', '100', '80'],
        ['Non-current assets:', '', ''], ['Investments', '50', '40'], ['Total non-current assets', '50', '40'],
        ['Total assets', '150', '120'],
        ['Current liabilities:', '', ''], ['Payables', '30', '20'], ['Total current liabilities', '30', '20'],
        ['Non-current liabilities:', '', ''], ['Debt', '50', '40'], ['Total non-current liabilities', '50', '40'],
        ['Total liabilities', '80', '60'],
        ['Shareholders’ equity:', '', ''], ['Common stock, 1,000 shares', '80', '65'], ['Accumulated deficit', '(10)', '(5)'],
        ['Total shareholders’ equity', '70', '60'], ['Total liabilities and shareholders’ equity', '150', '120'],
    ]
    return dict(rows=rows, row_kinds=['data' if row[1] else 'section' for row in rows],
                period_labels=['September 27, 2025', 'September 28, 2024'])


@pytest.mark.parametrize('label', ['Shareholders’ equity', "Stockholders' equity", 'Equity'])
def test_corporate_subtotals_and_equation_cover_both_dates(label):
    table = corporate()
    for row in table['rows']:
        row[0] = row[0].replace('Shareholders’ equity', label).replace('shareholders’ equity', label.lower())
    checks = audit_balance_grid(table)
    assert len(checks) == 18
    assert all(c['Status'] == 'PASS' for c in checks)


def test_detail_mistake_fails_even_when_printed_totals_balance():
    table = corporate()
    table['rows'][1][1] = '63'
    checks = audit_balance_grid(table)
    assert checks[0]['Status'] == 'FAIL'
    assert checks[8]['Status'] == 'PASS'
    assert all(c['Status'] == 'PASS' for c in checks[9:])


@pytest.mark.parametrize('mutation', ['blank', 'dash', 'missing_heading', 'duplicate_heading', 'nested_total', 'duplicate_equity'])
def test_ambiguous_or_unavailable_components_never_pass(mutation):
    table = corporate()
    if mutation in {'blank', 'dash'}:
        table['rows'][1][1] = '' if mutation == 'blank' else '—'
    elif mutation == 'missing_heading':
        table['rows'][0][0] = 'Unrecognized heading'
    elif mutation == 'duplicate_heading':
        table['rows'].append(deepcopy(table['rows'][0]))
        table['row_kinds'].append('section')
    elif mutation == 'nested_total':
        table['rows'][1][0] = 'Total cash'
    else:
        table['rows'].append(['Total net assets', '70', '60'])
        table['row_kinds'].append('data')
    checks = audit_balance_grid(table)
    assert checks[6 if mutation == 'duplicate_equity' else 0]['Status'] == 'NOT_TESTED'


def saved_cells():
    sheet = Workbook().active
    sheet['C5'] = 'Balance sheets'
    sheet['D8'], sheet['E8'] = 'September 27, 2025', 'September 28, 2024'
    # Repeated labels deliberately differ in value and source position.
    rows = [('Securities', 20, 10), ('Securities', 30, 40), ('Total assets', 50, 50)]
    reference = dict(id='balance', fields=[sheet['D8'].value, sheet['E8'].value], headings=[], rows=[])
    for i, (label, a, b) in enumerate(rows, 10):
        sheet.cell(i, 3, label)
        sheet.cell(i, 4, a)
        sheet.cell(i, 5, b)
        reference['rows'].append(dict(label=label, values=[a, b], source_lines=[i]))
    controls = [dict(check='Assets', period=p, source_input_lines=[10, 11], total_source_line=12)
                for p in reference['fields']]
    return sheet, reference, controls


def test_independent_arithmetic_uses_saved_cells_not_expected_values():
    sheet, ref, controls = saved_cells()
    assert all(c['status'] == 'PASS' for c in audit_saved_cells(sheet, ref, controls)[1])
    sheet['D10'] = 25  # The reference and its totals remain unchanged.
    checks = audit_saved_cells(sheet, ref, controls)[1]
    assert checks[0]['status'] == 'FAIL' and checks[0]['difference'] == '-5'
    assert checks[0]['input_cells'] == ['D10', 'D11']
    assert checks[1]['status'] == 'PASS'


@pytest.mark.parametrize('mutation', ['blank', 'formula', 'duplicate', 'shift', 'swap_dates', 'numeric_date'])
def test_independent_checks_refuse_unavailable_or_wrong_period_cells(mutation):
    sheet, ref, controls = saved_cells()
    if mutation == 'blank':
        sheet['D10'] = None
    elif mutation == 'formula':
        sheet['D10'] = '=10+10'
    elif mutation == 'duplicate':
        sheet['C13'], sheet['D13'], sheet['E13'] = 'Securities', 20, 10
    elif mutation == 'shift':
        sheet['F10'], sheet['D10'] = sheet['D10'].value, None
    elif mutation == 'swap_dates':
        sheet['D8'], sheet['E8'] = sheet['E8'].value, sheet['D8'].value
    else:
        sheet['D8'] = 2025
    checks = audit_saved_cells(sheet, ref, controls)[1]
    assert checks[0]['status'] == 'NOT_TESTED'


def test_companion_reference_must_be_manifest_registered_and_unchanged(tmp_path):
    files = {'source.pdf': b'source', 'before.xlsx': b'before', 'source-controls.json': b'[]'}
    files['expected.json'] = json.dumps({'source_sha256': hashlib.sha256(files['source.pdf']).hexdigest()}).encode()
    manifest = {'files': {name: {'sha256': hashlib.sha256(data).hexdigest()} for name, data in files.items()}}
    for name, data in files.items():
        (tmp_path / name).write_bytes(data)
    (tmp_path / 'manifest.json').write_text(json.dumps(manifest))
    load_reference(tmp_path, required_files=('source-controls.json',))
    (tmp_path / 'source-controls.json').write_text('[1]')
    with pytest.raises(ValueError, match='integrity'):
        load_reference(tmp_path, required_files=('source-controls.json',))
    del manifest['files']['source-controls.json']
    (tmp_path / 'manifest.json').write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match='register'):
        load_reference(tmp_path, required_files=('source-controls.json',))
