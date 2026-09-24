from copy import deepcopy
from types import SimpleNamespace

from openpyxl import load_workbook
import pandas as pd
import pytest

from excel_writer import _write_statement_sheet
from financial_statement_extract.equity_layout import EquityLayoutError, equity_component_grid, recover_pdf_equity_tables
from models import StatementPage, StatementPagePlan


def example():
    words = []
    def line(y, items):
        words.extend(dict(text=text, x0=x, x1=x+width, top=y) for x, width, text in items)
    line(76, [(310, 42, 'Common Stock')])
    line(83, [(400, 50, 'Paid in Capital in'), (478, 35, 'Distributable')])
    line(91, [(282, 30, 'Shares'), (344, 42, 'Par Amount'), (405, 45, 'Excess of Par'),
              (470, 50, 'Earnings (Losses)'), (531, 55, 'Total Net Assets')])
    labels = ['Balance at September 30, 2025', 'Issuance of common stock', 'Issuance of common stock',
              'Total increase for the nine months ended June 30, 2026', 'Balance at June 30, 2026']
    values = [['100', '$ 1', '$ 20', '$ (5)', '$ 16'], ['10', '—', '2', '—', '2'],
              ['10', '—', '2', '—', '2'], ['20', '—', '4', '—', '4'], ['120', '$ 1', '$ 24', '$ (5)', '$ 20']]
    for i, (label, vector) in enumerate(zip(labels, values)):
        line(102 + i*12, [(10, 230, label)] + [(x, 25, value) for x, value in zip([280, 360, 420, 480, 548], vector)])
    # An earlier, narrow parent-header table must not become the column schema.
    header = SimpleNamespace(rows=[SimpleNamespace(cells=[(260, 70, 395, 88)])], bbox=(260, 70, 395, 88))
    boundaries = [0, 260, 330, 395, 460, 525, 590]
    cells = [(left, 100, right, 155) for left, right in zip(boundaries, boundaries[1:])]
    body = SimpleNamespace(rows=[SimpleNamespace(cells=cells)], bbox=(0, 100, 590, 155))
    text = '\n'.join(' '.join([label, *vector]) for label, vector in zip(labels, values))
    page = SimpleNamespace(width=612, extract_words=lambda **kwargs: deepcopy(words), extract_text=lambda: text,
                           find_tables=lambda: [header, body])
    return page, words, labels, values


def test_geometry_uses_body_table_and_recovers_parent_and_leaf_headers():
    page, _, labels, values = example()
    grid = equity_component_grid(page, page.find_tables())
    assert grid['rows'][0] == ['', 'Common Stock', '', '', '', '']
    assert grid['component_labels'] == ['Shares', 'Par Amount', 'Paid in Capital in Excess of Par',
                                         'Distributable Earnings (Losses)', 'Total Net Assets']
    assert grid['spans'] == [(0, 1, 1, 2)]
    assert grid['rows'][2:] == [[label, *vector] for label, vector in zip(labels, values)]
    assert grid['rows'].count(['Issuance of common stock', '10', '—', '2', '—', '2']) == 2
    assert grid['period_sections'] == [dict(start=2, end=7, caption='nine months ended June 30, 2026', caption_source_row=5)]
    assert len(grid['row_sources']) == len(grid['rows'])


def test_missing_component_header_fails_instead_of_inventing_a_name():
    page, words, _, _ = example()
    words[:] = [word for word in words if word['text'] != 'Par Amount']
    with pytest.raises(EquityLayoutError, match='source header'):
        equity_component_grid(page, page.find_tables())


def test_two_amounts_in_one_component_are_not_concatenated():
    page, words, _, _ = example()
    next(word for word in words if word['text'] == '$ 20')['text'] = '20 30'
    with pytest.raises(EquityLayoutError, match='ambiguous numeric'):
        equity_component_grid(page, page.find_tables())


def test_missing_component_cell_is_not_shifted_into_another_column():
    page, words, _, _ = example()
    words[:] = [word for word in words if word['text'] != '$ 20']
    with pytest.raises(EquityLayoutError, match='missing component'):
        equity_component_grid(page, page.find_tables())


def test_source_text_disagreement_blocks_recovery():
    page, _, _, _ = example()
    text = page.extract_text().replace('$ 20', '$ 21')
    page.extract_text = lambda: text
    with pytest.raises(EquityLayoutError, match='complete source text'):
        equity_component_grid(page, page.find_tables())


def test_inconsistent_body_column_geometry_is_rejected():
    page, _, _, _ = example()
    found = page.find_tables()
    other = deepcopy(found[-1])
    other.rows[0].cells[1] = (260, 160, 325, 175)
    found.append(other)
    with pytest.raises(EquityLayoutError, match='disagree'):
        equity_component_grid(page, found)


def test_ambiguous_page_ownership_keeps_originals_with_a_review_warning():
    plan = StatementPagePlan('unused.pdf', 'manual', (1,), (1,), (1,),
                             (StatementPage(1, 'data', 'PartnersCapital', statement_types=('PartnersCapital', 'IncomeStatement')),))
    original = [dict(source_page=1, statement_type='PartnersCapital')]
    tables, issues = recover_pdf_equity_tables('unused.pdf', original, {}, plan)
    assert tables is original
    assert issues[0]['Status'] == 'WARN' and 'ambiguous' in issues[0]['Detail']


@pytest.mark.parametrize('conflict', [False, True])
def test_pipeline_recovery_preserves_raw_grids_and_warns_on_amount_conflict(monkeypatch, conflict):
    page, _, labels, vectors = example()
    class PDF:
        pages = [page]
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
    monkeypatch.setattr('financial_statement_extract.equity_layout.pdfplumber.open', lambda _: PDF())
    originals = [dict(table_id='raw', source_page=1, source_order=0, statement_type='PartnersCapital',
                      flavor='pdfplumber', rows=[[label, *vector] for label, vector in zip(labels, vectors)])]
    if conflict:
        originals[0]['rows'][0][-1] = '$ 17'
    saved = deepcopy(originals)
    statements = {'PartnersCapital': pd.DataFrame()}
    statements['PartnersCapital'].attrs['source_tables'] = originals
    plan = StatementPagePlan('file.pdf', 'manual', (1,), (1,), (1,),
                             (StatementPage(1, 'data', 'PartnersCapital', 'Changes in Net Assets'),))
    tables, issues = recover_pdf_equity_tables('file.pdf', originals, statements, plan)
    assert originals == saved
    if conflict:
        assert tables == originals
        assert statements['PartnersCapital'].attrs['source_tables'] == originals
        assert issues[0]['Status'] == 'WARN'
    else:
        assert tables[0]['flavor'] == 'equity_geometry'
        assert tables[0]['source_table_ids'] == ['raw']
        assert issues[0]['Status'] == 'INFO'


def test_saved_equity_uses_stable_columns_full_headers_totals_and_no_freeze_panes(tmp_path):
    page, _, _, _ = example()
    grid = equity_component_grid(page, page.find_tables())
    grid.update(flavor='equity_geometry', source_context=['(In thousands, except share data)'])
    frame = pd.DataFrame()
    frame.attrs.update(source_tables=[grid], company_name='Example Fund',
                       statement_title='Consolidated Statements of Changes in Net Assets (unaudited)')
    path = tmp_path / 'equity.xlsx'
    with pd.ExcelWriter(path, engine='xlsxwriter') as writer:
        _write_statement_sheet(writer, 'PartnersCapital', frame)
    book = load_workbook(path)
    sheet = book['Changes in Net Assets']
    assert sheet.freeze_panes is None
    assert sheet.max_column == 8
    assert sheet['D9'].value == 'Common Stock'
    assert not sheet.merged_cells.ranges
    assert all(sheet[cell].alignment.horizontal == 'centerContinuous' for cell in ('D9', 'E9'))
    assert [c.value for c in sheet[10][2:]][1:] == grid['component_labels']
    data = [r for r in sheet.iter_rows(min_col=3) if r[0].value == 'Issuance of common stock']
    assert len(data) == 2
    assert [c.value for c in data[0]] == ['Issuance of common stock', 10, '—', 2, '—', 2]
    total = next(r for r in sheet.iter_rows(min_col=3) if str(r[0].value).startswith('Total increase'))
    assert all(c.font.bold and c.border.top.style for c in total)
    assert not any(c.data_type == 'f' for row in sheet.iter_rows(min_col=3) for c in row)
    book.close()
