"""Completeness must describe the exported grid, even when recovery fails."""
from copy import deepcopy
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

from excel_writer import write_extraction_workbook
from models import ExtractionResult
from financial_statement_extract.ui.results import summarize_result
from financial_statement_extract.workbook_validation import (
    BALANCE_COMPLETENESS, audit_balance_completeness, audit_selected_page_coverage,
    balance_completeness_status,
)


def sample(flavor='balance_geometry'):
    rows = [['ASSETS:', '', ''], ['Cash', '$ 60', '$ 50'], ['Other assets', '40', '40'],
            ['Total assets', '100', '90'], ['LIABILITIES:', '', ''],
            ['Total liabilities', '60', '50'], ['Shareholders’ equity:', '', ''],
            ['Common stock, $0.01 par value; 1,000 shares authorized', '40', '40'],
            ['Total shareholders’ equity', '40', '40'],
            ['Total liabilities and shareholders’ equity', '100', '90']]
    lines = ['Company', 'Balance Sheet', '2026 2025', *[' '.join(row).strip() for row in rows],
             'See accompanying Notes to Consolidated Financial Statements.']
    # The monetary row has a multiline label containing unrelated share counts.
    lines[10:11] = ['Common stock, $0.01 par value;', '1,000 shares authorized 40 40']
    table = dict(table_id='balance', source_page=1, flavor=flavor, rows=rows,
                 row_kinds=['data' if row[1] else 'section' for row in rows],
                 row_sources=[[{'line_no': i + 1}] for i in range(len(rows))],
                 period_labels=['2026', '2025'])
    frame = pd.DataFrame()
    frame.attrs.update(source_tables=[table], statement_title='Balance Sheet', company_name='Company')
    page = SimpleNamespace(page=1, statement_type='BalanceSheet', statement_types=())
    return ExtractionResult(
        metadata={'source_file': 'sample.pdf', 'pages_selected': '1'},
        raw_text_rows=[dict(source_page=1, line_no=i+1, raw_text=line) for i, line in enumerate(lines)],
        raw_table_cells=[], normalized_table_cells=[], extraction_audit_rows=[],
        statements={'BalanceSheet': frame}, parsed_cells=[], financial_audit_rows=[], unmapped_rows=[],
        page_plan=SimpleNamespace(pages=[page], selected_pages=[1], warnings=()),
    )


def check(result):
    return audit_balance_completeness(result, result.statements.get('BalanceSheet'), 1)


@pytest.mark.parametrize('flavor', ['balance_geometry', 'stream'])
def test_complete_wrapped_balance_passes_even_without_recovery(flavor):
    result = sample(flavor)
    assert check(result)['Status'] == 'PASS'
    assert 'content coverage only' in check(result)['Detail']


@pytest.mark.parametrize('flavor', ['balance_geometry', 'stream'])
def test_assets_only_export_fails_despite_full_analytical_frame(flavor):
    result = sample(flavor)
    frame = result.statements['BalanceSheet']
    frame['RawItem'] = ['Total liabilities', 'Total shareholders’ equity']
    frame.attrs['source_tables'][0]['rows'] = frame.attrs['source_tables'][0]['rows'][:4]
    outcome = check(result)
    assert outcome['Status'] == 'FAIL'
    assert 'equity/net assets, liabilities' in outcome['Detail']


@pytest.mark.parametrize('mutation', ['remove_detail', 'duplicate_replaces_detail', 'remove_heading', 'lose_share_count'])
def test_missing_content_fails_even_when_totals_still_match(mutation):
    result = sample()
    rows = result.statements['BalanceSheet'].attrs['source_tables'][0]['rows']
    if mutation == 'remove_detail':
        rows.pop(2)
    elif mutation == 'duplicate_replaces_detail':
        rows[2] = deepcopy(rows[1])
    elif mutation == 'remove_heading':
        rows.pop(4)
    else:
        rows[7][0] = 'Common stock, $0.01 par value; shares authorized'
    assert check(result)['Status'] == 'FAIL'


def test_currency_columns_whitespace_and_explicit_zero_are_preserved():
    result = sample('stream')
    rows = result.statements['BalanceSheet'].attrs['source_tables'][0]['rows']
    rows[1] = ['Cash', '$', '60', '$', '50']
    rows[7][0] = 'Common stock,\n$ 0.01 par value; 1,000 shares authorized'
    rows.append(['Other', 0, '—'])
    result.raw_text_rows.insert(-1, dict(source_page=1, raw_text='Other 0 —'))
    assert check(result)['Status'] == 'PASS'
    rows[-1][1] = ''
    assert check(result)['Status'] == 'FAIL'


@pytest.mark.parametrize('mutation', ['no_text', 'no_end', 'duplicate_start', 'no_source_equity'])
def test_unbounded_or_unsupported_source_is_unverified_never_passed(mutation):
    result = sample()
    if mutation == 'no_text':
        result.raw_text_rows = []
    elif mutation == 'no_end':
        result.raw_text_rows.pop()
    elif mutation == 'duplicate_start':
        result.raw_text_rows.insert(4, dict(source_page=1, raw_text='Assets'))
    else:
        result.raw_text_rows = [row for row in result.raw_text_rows if 'equity' not in row['raw_text']]
    assert check(result)['Status'] == 'NOT_TESTED'


def test_missing_entire_balance_sheet_has_explicit_incomplete_status():
    result = sample()
    result.statements = {}
    result.extraction_audit_rows = audit_selected_page_coverage(result)
    assert balance_completeness_status(result.extraction_audit_rows) == 'FAIL'
    assert 'INCOMPLETE BALANCE SHEET' in summarize_result(result).headline


@pytest.mark.parametrize('status', ['FAIL', 'NOT_TESTED', 'PASS'])
@pytest.mark.parametrize('flavor', ['balance_geometry', 'stream'])
def test_workbook_and_app_show_same_explicit_state(tmp_path, status, flavor):
    result = sample(flavor)
    if status == 'FAIL':
        result.statements['BalanceSheet'].attrs['source_tables'][0]['rows'].pop()
        # Keep recovered structural metadata aligned for this test.
        for field in ('row_kinds', 'row_sources'):
            result.statements['BalanceSheet'].attrs['source_tables'][0][field].pop()
    elif status == 'NOT_TESTED':
        result.raw_text_rows = []
    result.extraction_audit_rows = audit_selected_page_coverage(result)
    assert balance_completeness_status(result.extraction_audit_rows) == status
    summary = summarize_result(result)
    book = load_workbook(write_extraction_workbook(result, tmp_path / 'result.xlsx'))
    sheet = book['Balance Sheet']
    if status == 'FAIL':
        assert 'INCOMPLETE BALANCE SHEET' in summary.headline
        assert 'INCOMPLETE BALANCE SHEET' in sheet['C5'].value
        assert 'INCOMPLETE BALANCE SHEET' in book['Review']['C4'].value
    elif status == 'NOT_TESTED':
        assert 'completeness unverified' in summary.headline
        assert 'COMPLETENESS UNVERIFIED' in sheet['C5'].value
        assert 'completeness unverified' in book['Review']['C4'].value
    else:
        assert sheet['C5'].value == 'Balance Sheet'
        assert 'INCOMPLETE' not in summary.headline
        assert any(row[3].value == BALANCE_COMPLETENESS for row in book['Review'])
    for tab in book:
        assert not tab.merged_cells.ranges and tab.freeze_panes is None
        assert all(cell.value is None for row in tab.iter_rows(max_row=3) for cell in row)
        assert all(cell.value is None for row in tab.iter_rows(max_col=2) for cell in row)
        assert all(not dim.hidden for dim in [*tab.row_dimensions.values(), *tab.column_dimensions.values()])
        assert tab.column_dimensions['A'].width == pytest.approx(1.7109375)
    book.close()


def test_failures_take_precedence_over_passed_pages_or_arithmetic():
    rows = [dict(Check=BALANCE_COMPLETENESS, Status=status) for status in ['PASS', 'NOT_TESTED', 'FAIL']]
    rows.append(dict(Check='Balance arithmetic', Status='PASS'))
    assert balance_completeness_status(rows) == 'FAIL'
