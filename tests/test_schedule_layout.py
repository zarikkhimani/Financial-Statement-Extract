from copy import deepcopy
from types import SimpleNamespace

from openpyxl import load_workbook
import pandas as pd
import pytest
import xlsxwriter

import extractors
from excel_writer import _write_source_grids
from financial_statement_extract.schedule_layout import merge_overlapping_schedule_tables, schedule_display_grid
from financial_statement_extract.table_layout import organize_source_tables, selected_grids
from models import StatementPage, StatementPagePlan


def grid(name, rows, positions=None, page=8):
    positions = positions or [700 - i * 12 for i in range(len(rows))]
    return dict(table_id=name, source_page=page, source_order=1, flavor="stream",
                statement_type="ScheduleOfInvestments", source_title="Schedule of Investments",
                rows=rows, row_bounds=[(y, y - 10) for y in positions],
                column_bounds=[(c * 100, (c + 1) * 100) for c in range(len(rows[0]))],
                spans=[], source_context=[])


def test_triplicate_and_headerless_subset_keep_real_identical_investments():
    rows = [["Investment", "Value"], ["API", "100"], ["API", "100"], ["Total", "200"]]
    tables = [grid("full", rows), grid("copy", rows), grid("subset", rows[1:], [688, 676, 664])]
    original = deepcopy(tables)
    # A cropped label boundary and a slightly different top row are normal.
    tables[2]["column_bounds"][0] = (6, 100)
    tables[2]["row_bounds"][0] = (687, 678)
    original = deepcopy(tables)
    merged, issues = merge_overlapping_schedule_tables(tables)
    assert len(merged) == 1
    assert merged[0]["rows"] == rows
    assert merged[0]["source_table_ids"] == ["full", "copy", "subset"]
    assert [len(sources) for sources in merged[0]["row_sources"]] == [2, 3, 3, 3]
    assert merged[0]["row_sources"][2][-1] == {"table_id": "subset", "row_index": 1}
    assert all(issue["Status"] == "INFO" for issue in issues)
    assert tables == original


def test_partial_overlap_can_bridge_two_disjoint_tables_without_losing_boundary_rows():
    tables = [grid("top", [["A", "1"], ["B", "2"]], [700, 688]),
              grid("bottom", [["D", "4"], ["E", "5"]], [664, 652]),
              grid("bridge", [["B", "2"], ["C", "3"], ["D", "4"]], [688, 676, 664])]
    merged, _ = merge_overlapping_schedule_tables(tables)
    assert len(merged) == 1
    assert merged[0]["rows"] == [[letter, str(i)] for i, letter in enumerate("ABCDE", 1)]
    assert [len(sources) for sources in merged[0]["row_sources"]] == [1, 2, 1, 2, 1]


@pytest.mark.parametrize("difference", ["page", "position", "horizontal", "html", "statement"])
def test_equal_values_in_distinct_source_locations_or_non_pdf_schedules_survive(difference):
    first = grid("a", [["API", "100"]])
    second = grid("b", [["API", "100"]])
    if difference == "page":
        second["source_page"] = 9
    elif difference == "position":
        second["row_bounds"] = [(688, 678)]
    elif difference == "horizontal":
        second["column_bounds"] = [(300, 400), (400, 500)]
    elif difference == "html":
        first["flavor"] = second["flavor"] = "html"
    else:
        first["statement_type"] = second["statement_type"] = "IncomeStatement"
    merged, issues = merge_overlapping_schedule_tables([first, second])
    assert len(merged) == 2
    assert not issues


@pytest.mark.parametrize("missing", ["row_bounds", "column_bounds", "source_page"])
def test_missing_geometry_warns_and_never_deduplicates_by_content(missing):
    tables = [grid("a", [["API", "100"]]), grid("b", [["API", "100"]])]
    for table in tables:
        table.pop(missing)
    merged, issues = merge_overlapping_schedule_tables(tables)
    assert len(merged) == 2
    assert issues[0]["Status"] == "WARN"


def test_overlapping_candidates_with_incompatible_columns_warn_and_remain_separate():
    first = grid("a", [["API", "100"]])
    second = grid("b", [["API", "100"]])
    second["column_bounds"] = [(0, 80), (80, 200)]
    merged, issues = merge_overlapping_schedule_tables([first, second])
    assert len(merged) == 2
    assert "different column boundaries" in issues[0]["Detail"]


def test_conflicting_amounts_at_same_position_fail_clearly():
    with pytest.raises(ValueError, match="Conflicting overlapping schedule"):
        merge_overlapping_schedule_tables([grid("a", [["API", "100"]]), grid("b", [["API", "101"]])])


def test_ambiguous_geometry_fails_instead_of_removing_genuine_repeats():
    first = grid("a", [["API", "100"], ["API", "100"]], [700, 699])
    with pytest.raises(ValueError, match="Ambiguous overlapping"):
        merge_overlapping_schedule_tables([first, grid("b", [["API", "100"]])])


def test_pik_suffix_moves_only_with_proven_headers_rate_and_date():
    table = grid("a", [["Investment", "Interest", "Maturity"],
                       ["API", "6.14%\n cash/\n2.88%", "PIK\n07/2032"],
                       ["Fixed", "13.00%\nPIK", "04/2030"],
                       ["Unknown", "See note", "PIK\n07/2032"],
                       ["Ambiguous", "6.14% cash/2.88%", "PIK date pending"]])
    original = deepcopy(table)
    display = schedule_display_grid(table)
    assert display["rows"][1] == ["API", "6.14%\n cash/\n2.88% PIK", "07/2032"]
    assert display["rows"][2:] == table["rows"][2:]
    assert table == original
    table["rows"][0][1] = "Unrelated column"
    assert schedule_display_grid(table)["rows"] == table["rows"]


def test_extractor_keeps_geometry_and_all_raw_candidates(monkeypatch):
    table = SimpleNamespace(df=pd.DataFrame([["API", "100"]]),
                            rows=[(700, 690)], cols=[(0, 100), (100, 200)],
                            parsing_report={"accuracy": 99, "whitespace": 0, "order": 1})
    monkeypatch.setattr(extractors, "get_page_count", lambda _: 9)
    monkeypatch.setattr(extractors.camelot, "read_pdf",
                        lambda *args, **kwargs: [table, table, table] if kwargs["flavor"] == "stream" else [])
    cells, _, _ = extractors.extract_tables_camelot("unused.pdf", "8")
    assert len(cells) == 6
    assert all(cell["row_bounds"] == [700, 690] for cell in cells)
    assert {tuple(cell["column_bounds"]) for cell in cells} == {(0, 100), (100, 200)}
    assert len(selected_grids(cells)) == 3


def test_organized_schedule_exports_date_and_clean_maturity_without_changing_raw_cells(tmp_path):
    rows = [["Investment", "Interest", "Maturity", "Value"],
            ["API", "6.14% cash/2.88%", "PIK\n07/2032", "100"]]
    cells = [dict(table_id=name, selected=True, source_page=8, flavor="stream",
                  row_index=r, column_index=c, raw_text=value, row_bounds=(700-r*12, 690-r*12),
                  column_bounds=(c*100, (c+1)*100))
             for name in ("a", "b", "c") for r, row in enumerate(rows) for c, value in enumerate(row)]
    original = deepcopy(cells)
    page = StatementPage(8, "data", "ScheduleOfInvestments", "Schedule of Investments")
    plan = StatementPagePlan("source.pdf", "manual", (8,), (8,), (8,), (page,))
    frames = {}
    tables, _ = organize_source_tables(cells, [dict(source_page=8, raw_text="June 30, 2026")], frames, plan)
    assert len(tables) == 1
    assert tables[0]["rows"] == rows
    assert tables[0]["source_context"] == ["June 30, 2026"]
    assert cells == original
    path = tmp_path / "schedule.xlsx"
    with xlsxwriter.Workbook(path) as book:
        _write_source_grids(book, book.add_worksheet("Schedule"), frames["ScheduleOfInvestments"])
    book = load_workbook(path)
    sheet = book["Schedule"]
    assert sheet["C8"].value == "June 30, 2026"
    assert sheet["D10"].value == "6.14% cash/2.88% PIK"
    assert sheet["E10"].value == "07/2032"
    assert sheet["F10"].value == 100
    assert len([row for row in sheet.iter_rows(min_col=3, values_only=True) if row[0] == "API"]) == 1
    book.close()
