from pathlib import Path
from types import SimpleNamespace

from bs4 import BeautifulSoup
from openpyxl import load_workbook
import pytest

from financial_statement_extract.pdf_layout import equity_grid
from financial_statement_extract.table_layout import covered_pages, equity_display_grid, organize_source_tables, selected_grids
from html_extractor import _table_grid
from models import StatementPage, StatementPagePlan
from pipeline import extract_html_to_workbook


EQUITY = """<html><body><h2>Statements of Stockholders’ Equity</h2><table>
<tr><th rowspan="2"></th><th colspan="2">Common Stock</th><th rowspan="2">Retained earnings (deficit)</th></tr>
<tr><th>Shares</th><th>Amount</th></tr>
<tr><td>Balance at December 31, 2025</td><td>100</td><td>$1.00</td><td>20</td></tr>
<tr><td>Net increase in net assets from operations (3)</td><td>—</td><td>N/A</td><td>5</td></tr>
<tr><td>Balance at June 30, 2026</td><td>100</td><td>$1.00</td><td>25</td></tr>
</table><h2>Schedule of Investments</h2><table>
<tr><th colspan="4">June 30, 2026</th></tr>
<tr><th>Investment</th><th>Asset</th><th>Cost</th><th>Fair Value</th></tr>
<tr><td>Company A (1)</td><td>Preferred stock</td><td>2</td><td>3</td></tr>
<tr><td>Company B</td><td>Debt</td><td>4</td><td>5</td></tr>
</table><table>
<tr><th colspan="3">December 31, 2025</th></tr>
<tr><th>Investment</th><th>Geography</th><th>% of net assets</th></tr>
<tr><td>Company A (1)</td><td>Europe</td><td>12.50%</td></tr>
<tr><td>Company B</td><td>Americas</td><td>0%</td></tr>
</table></body></html>"""


def test_html_grid_keeps_spans_and_column_positions():
    table = BeautifulSoup(EQUITY, "lxml").find("table")
    _, rows, spans = _table_grid(table)
    assert rows[:2] == [["", "Common Stock", "", "Retained earnings (deficit)"], ["", "Shares", "Amount", ""]]
    assert spans[(0, 1)] == (1, 2)
    assert spans[(0, 3)] == (2, 1)


def test_equity_and_incompatible_schedule_schemas_export_without_relabeling(tmp_path):
    source = tmp_path / "equity.html"
    source.write_text(EQUITY, encoding="utf-8")
    result, output = extract_html_to_workbook(str(source), output_dir=str(tmp_path))
    equity = result.statements["StockholdersEquityStatement"]
    assert equity.attrs["period_labels"] == []  # row dates are not value-column headers
    schedules = [b for b in result.statement_tables if b["statement_type"] == "ScheduleOfInvestments"]
    assert len(schedules) == 2
    assert len(schedules[0]["rows"][0]) == 4
    assert len(schedules[1]["rows"][0]) == 3
    book = load_workbook(output)
    eq = next(sheet for sheet in book if sheet["C5"].value == "Statements of Stockholders’ Equity")
    assert eq["D8"].value == "Common Stock"
    assert not eq.merged_cells.ranges
    assert all(eq[cell].alignment.horizontal == 'centerContinuous' for cell in ('D8', 'E8'))
    assert eq["C10"].value == "Balance at December 31, 2025"
    assert eq["E10"].value == 1
    assert '"$"' in eq["E10"].number_format
    assert eq["C11"].value == "Net increase in net assets from operations (3)"
    assert eq["D11"].value == "—"
    assert eq["E11"].value == "N/A"
    schedule = book["Schedule of Investments"]
    cells = [c for row in schedule for c in row]
    assert {"Preferred stock", "Geography", "Europe", "December 31, 2025"} <= {c.value for c in cells}
    assert any(c.value == 0.125 and "0.00%" in c.number_format for c in cells)
    assert not any(c.data_type == "f" for c in cells)
    book.close()


def cell_rows(rows, page=1):
    return [dict(table_id="test", selected=True, source_page=page, row_index=r, column_index=c, raw_text=text)
            for r, row in enumerate(rows) for c, text in enumerate(row)]


def test_coverage_counts_duplicate_values_and_requires_period_headers():
    source = [dict(source_page=1, raw_text="2026 2025\nFirst 100 20\nSecond 100 20")]
    assert covered_pages(cell_rows([["2026", "2025"], ["First", "100", "20"]]), source) == set()
    assert covered_pages(cell_rows([["2026", "2025"], ["First", "100", "20"], ["Second", "100", "20"]]), source) == {1}
    many = [dict(source_page=1, raw_text="2026 2025\n" + "\n".join(["Line 100 20"] * 100))]
    assert covered_pages(cell_rows([["Line", "100", "20"]] * 100), many) == set()


def test_engine_execution_order_does_not_reorder_source_pages():
    cells = []
    for page in [6, 10, 7, 8, 9]:
        records = cell_rows([["Company", "100", "20"]], page)
        for cell in records:
            cell["table_id"] = str(page)
        cells.extend(records)
    assert [block["source_page"] for block in selected_grids(cells)] == [6, 7, 8, 9, 10]


def test_detected_equity_without_a_layout_fails_before_writing():
    page = StatementPage(1, "data", "StockholdersEquityStatement", "Statements of Stockholders' Equity")
    plan = StatementPagePlan("source.pdf", "manual", (1,), (1,), (1,), (page,))
    with pytest.raises(ValueError, match="column layout"):
        organize_source_tables([], [], {}, plan)


def test_equity_html_layout_subdivisions_collapse_to_source_components():
    block = dict(flavor="html", statement_type="StockholdersEquityStatement", rows=[
        ["", "", "Common Stock", "", "", "", "Earnings", ""],
        ["", "", "Shares", "", "Amount", "", "", ""],
        ["Balance at December 31, 2025", "", "100", "", "$", "1", "$", "20"],
        ["Balance at June 30, 2026", "", "100", "", "$", "1", "$", "25"],
    ], spans=[(0, 2, 1, 4), (0, 6, 2, 2), (1, 2, 1, 2), (1, 4, 1, 2), (2, 0, 1, 2)])
    displayed = equity_display_grid(block)
    assert displayed["rows"][2] == ["Balance at December 31, 2025", "100", "$ 1", "$ 20"]
    assert (0, 1, 1, 2) in displayed["spans"]
    assert len(block["rows"][2]) == 8  # raw provenance is unchanged


def test_numeric_cells_inside_source_colspans_stay_numeric(tmp_path):
    html = EQUITY.replace("<td>100</td>", '<td colspan="2">100</td>')
    source = tmp_path / "merged.html"
    source.write_text(html, encoding="utf-8")
    _, output = extract_html_to_workbook(str(source), output_dir=str(tmp_path))
    book = load_workbook(output)
    equity = next(s for s in book if s["C5"].value == "Statements of Stockholders’ Equity")
    assert any(c.value == 100 and c.data_type == "n" for row in equity for c in row)
    book.close()


def test_multiple_numeric_tokens_in_a_cell_are_not_concatenated(tmp_path):
    source = tmp_path / "ambiguous.html"
    source.write_text(EQUITY.replace("<td>100</td>", "<td>100 200</td>"), encoding="utf-8")
    result, output = extract_html_to_workbook(str(source), output_dir=str(tmp_path))
    book = load_workbook(output)
    cells = [c for sheet in book for row in sheet.iter_rows(min_col=3) for c in row]
    assert any(c.value == "100 200" and c.data_type == "s" for c in cells)
    assert not any(c.value == 100200 for c in cells)
    assert any("ambiguous numeric" in row["Detail"] for row in result.extraction_audit_rows)
    assert book.sheetnames[-1] == "Review"
    book.close()


def test_pdf_equity_geometry_includes_open_header_and_row_between_table_boxes():
    words = []
    def line(y, *items):
        words.extend(dict(text=text, x0=x, x1=x + width, top=y) for x, width, text in items)
    line(50, (220, 40, "(unaudited)"))
    line(70, (310, 30, "Shares"), (370, 30, "Amount"), (470, 30, "Earnings"))
    line(100, (30, 180, "Balance at December 31, 2025"), (310, 20, "100"), (380, 20, "1"), (480, 20, "20"))
    line(120, (30, 180, "Balance at March 31, 2026"), (310, 20, "100"), (380, 20, "1"), (480, 20, "25"))
    line(140, (30, 180, "Balance at June 30, 2026"), (310, 20, "100"), (380, 20, "1"), (480, 20, "30"))
    cells = [(25, 95, 300, 110), (300, 95, 345, 110), (345, 95, 420, 110), (420, 95, 530, 110)]
    tables = [SimpleNamespace(rows=[SimpleNamespace(cells=cells)], bbox=(25, 95, 530, 110)),
              SimpleNamespace(bbox=(25, 135, 530, 150))]
    page = SimpleNamespace(width=612, extract_words=lambda **kwargs: words)
    grid = equity_grid(page, tables)
    assert grid[0] == ["", "Shares", "Amount", "Earnings"]
    assert grid[2] == ["Balance at March 31, 2026", "100", "1", "25"]


@pytest.mark.parametrize("strategy,expected_calls", [("adaptive", []), ("all", ["1"])])
def test_pdf_scheduler_skips_redundant_methods_only_with_coverage(tmp_path, monkeypatch, strategy, expected_calls):
    import pipeline
    source = tmp_path / "source.pdf"
    source.touch()
    rows = [dict(source_page=1, raw_text=line) for line in ["Income Statements", "2026 2025", "Revenue 100 90", "Net income 10 9"]]
    cells = cell_rows([["", "2026", "2025"], ["Revenue", "100", "90"], ["Net income", "10", "9"]])
    calls = []
    monkeypatch.setattr(pipeline, "get_page_count", lambda _: 1)
    monkeypatch.setattr(pipeline, "get_pdf_document_type", lambda _: "")
    monkeypatch.setattr(pipeline, "extract_page_text_pdfplumber", lambda *args: (rows, []))
    monkeypatch.setattr(pipeline, "extract_tables_pdfplumber", lambda *args: (cells, [], []))
    experimental = [dict(cells[0], flavor="pdfplumber_experimental",
                         table_id="EXPERIMENT", raw_text="=999+1", selected=True)]
    experimental_calls = []
    def experiment(path, pages):
        experimental_calls.append(pages)
        return experimental, [], [{"Check": "PDFPlumber Experimental table extraction",
                                   "Status": "ERROR", "Scope": "Page 1", "Detail": "Experiment diagnostic"}]
    monkeypatch.setattr(pipeline, "extract_tables_pdfplumber_experimental", experiment)
    monkeypatch.setattr(pipeline, "extract_tables_camelot", lambda _, pages: (calls.append(pages) or [], [], []))
    result, path = pipeline.extract_pdf_to_workbook(str(source), "1", output_dir=str(tmp_path), table_strategy=strategy)
    assert calls == expected_calls
    assert Path(path).exists()
    assert result.statement_tables[0]["rows"][1] == ["Revenue", "100", "90"]
    assert experimental_calls == ["1"]
    assert result.experimental_raw_table_cells == experimental
    assert all(cell.get("table_id") != "EXPERIMENT" for cell in result.raw_table_cells)
    assert all(row.get("Detail") != "Experiment diagnostic" for row in result.extraction_audit_rows)
    book = load_workbook(path)
    try:
        sheet = book["Experiential"]
        assert sheet["C8"].value == "=999+1"
        assert sheet["C8"].data_type == "s"
        assert any(cell.value == "Experiment diagnostic" for row in book["Review"].iter_rows(min_col=3) for cell in row)
    finally:
        book.close()
