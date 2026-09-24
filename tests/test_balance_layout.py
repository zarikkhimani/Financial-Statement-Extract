from copy import deepcopy
from types import SimpleNamespace

from openpyxl import load_workbook
import pandas as pd
import pytest

from excel_writer import _write_statement_sheet
from extractors import select_auto_statement_pages
from financial_statement_extract.balance_layout import BalanceLayoutError, balance_period_grid, recover_pdf_balance_tables
from models import StatementPage, StatementPagePlan
from structure import detect_statement_heading_type, split_statement_sections


def example():
    labels = ['Assets', 'Cash', 'Common stock, $0.001 par, 500,000 shares authorized,\n100 shares outstanding',
              'Total Assets']
    vectors = [['', ''], ['$ 1,000', '$ 900'], ['100', '90'], ['$ 1,100', '$ 990']]
    physical, extracted, words = [], [], []
    for i, (label, vector) in enumerate(zip(labels, vectors)):
        top = 100 + i * 24
        cells = [(0, top, 300, top+24), (300, top, 315, top+24), (315, top, 440, top+24),
                 (440, top, 450, top+24), (450, top, 465, top+24), (465, top, 590, top+24)]
        values = [label, '', '', '', '', '']
        for col, value in zip([1, 4], vector):
            if value.startswith('$'):
                values[col], values[col+1] = '$', value[2:]
            else:
                cells[col] = (cells[col][0], top, cells[col+1][2], top+24)
                cells[col+1] = None
                values[col] = value
        physical.append(SimpleNamespace(cells=cells))
        extracted.append(values)
        for value, cell in zip(values, cells):
            if value:
                words.append(dict(text=value, x0=cell[0]+1, x1=cell[2]-1, top=top+2, bottom=top+20))
    # A spanning audit-status header establishes the two major date columns.
    header = SimpleNamespace(cells=[(0, 80, 300, 100), (300, 80, 440, 100), None,
                                    (440, 80, 450, 100), (450, 80, 590, 100), None])
    physical.insert(0, header)
    extracted.insert(0, ['', '(unaudited)', '', '', '', ''])
    words.extend([dict(text='(unaudited)', x0=330, x1=400, top=82, bottom=94),
                  dict(text='June 30, 2026', x0=320, x1=430, top=66, bottom=76),
                  dict(text='September 30, 2025', x0=460, x1=580, top=66, bottom=76)])
    table = SimpleNamespace(rows=physical, bbox=(0, 80, 590, 196), extract=lambda: deepcopy(extracted))
    text = '\n'.join(['Example Company', 'Consolidated Statements of Financial Condition',
                      'June 30, 2026 September 30, 2025', '(unaudited)'] +
                     [' '.join([label, *vector]) for label, vector in zip(labels, vectors)] +
                     ['See Notes to Consolidated Financial Statements.', '3'])
    page = SimpleNamespace(width=612, find_tables=lambda: [table], extract_text=lambda: text,
                           extract_words=lambda **kwargs: deepcopy(words))
    return page, table, extracted, words


@pytest.mark.parametrize('title', ['Statement of Financial Condition', 'Consolidated Statements of Financial Condition'])
def test_financial_condition_heading_is_a_balance_sheet(title):
    assert detect_statement_heading_type(title) == 'BalanceSheet'
    text = f'{title}\nJune 30, 2026 September 30, 2025\nAssets\nCash 100 90\nTotal assets 100 90\nLiabilities 40 30\nNet assets 60 60'
    assert select_auto_statement_pages([text]) == [1]


def test_cash_reconciliation_reference_does_not_change_statement_ownership():
    texts = ['Consolidated Statements of Cash Flows', '2026 2025',
             'Consolidated Statements of Financial Condition that sum to the cash flow totals:', 'Cash 10 20']
    sections = split_statement_sections([dict(raw_text=t) for t in texts])
    assert list(sections) == ['CashFlowStatement']
    assert [r['raw_text'] for r in sections['CashFlowStatement']] == texts


def test_currency_cells_and_wrapped_stock_labels_keep_their_date_columns():
    page, _, _, _ = example()
    grid = balance_period_grid(page, page.find_tables())
    assert grid['period_labels'] == ['June 30, 2026', 'September 30, 2025']
    assert grid['period_notes'] == ['(unaudited)', '']
    assert grid['rows'][1] == ['Cash', '$ 1,000', '$ 900']
    assert grid['rows'][2] == ['Common stock, $0.001 par, 500,000 shares authorized, 100 shares outstanding', '100', '90']
    assert grid['row_kinds'] == ['section', 'data', 'data', 'data']
    assert len(grid['row_sources']) == 4


@pytest.mark.parametrize('mutation, message', [('date', 'source date'), ('row', 'complete source body'),
                                               ('value', 'source words'), ('crossing', 'crosses date columns')])
def test_uncertain_geometry_or_incomplete_content_is_rejected(mutation, message):
    page, table, extracted, words = example()
    if mutation == 'date':
        words[:] = [w for w in words if w['text'] != 'June 30, 2026']
    elif mutation == 'row':
        table.rows.pop()
        extracted.pop()
    elif mutation == 'value':
        extracted[2][2] = '1,001'
    else:
        table.rows[2].cells[2] = (315, 124, 470, 148)
    with pytest.raises(BalanceLayoutError, match=message):
        balance_period_grid(page, page.find_tables())


@pytest.mark.parametrize('failure', [False, True])
def test_recovery_preserves_raw_candidates_and_reports_failure(monkeypatch, failure):
    page, _, _, words = example()
    if failure:
        words[:] = [w for w in words if w['text'] != 'June 30, 2026']
    class PDF:
        pages = [page]
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
    monkeypatch.setattr('financial_statement_extract.balance_layout.pdfplumber.open', lambda _: PDF())
    original = [dict(table_id='raw', source_page=1, source_order=0, statement_type='BalanceSheet',
                     flavor='lattice', rows=[['Cash $ 1,000', '$ 900']])]
    saved = deepcopy(original)
    statements = {}
    plan = StatementPagePlan('file.pdf', 'manual', (1,), (1,), (1,),
                             (StatementPage(1, 'data', 'BalanceSheet', 'Consolidated Statements of Financial Condition'),))
    tables, issues = recover_pdf_balance_tables('file.pdf', original, statements, plan)
    assert original == saved
    assert issues[0]['Status'] == ('WARN' if failure else 'INFO')
    if failure:
        assert tables is original and not statements
        assert 'Original grids retained' in issues[0]['Detail']
    else:
        assert tables[0]['flavor'] == 'balance_geometry'
        assert tables[0]['source_table_ids'] == ['raw']
        assert statements['BalanceSheet'].attrs['source_tables'] == tables


def test_export_keeps_source_units_dates_full_labels_numeric_types_and_no_freeze(tmp_path):
    page, _, _, _ = example()
    grid = balance_period_grid(page, page.find_tables())
    grid.update(flavor='balance_geometry', source_context=['(In thousands, except share and per share data)'])
    frame = pd.DataFrame()
    frame.attrs.update(source_tables=[grid], company_name='Example Company',
                       statement_title='Consolidated Statements of Financial Condition')
    path = tmp_path / 'balance.xlsx'
    with pd.ExcelWriter(path, engine='xlsxwriter') as writer:
        _write_statement_sheet(writer, 'BalanceSheet', frame)
    book = load_workbook(path)
    sheet = book['Balance Sheet']
    assert sheet.freeze_panes is None and sheet.max_column == 5
    assert sheet['C7'].value == grid['source_context'][0]
    assert [sheet['D8'].value, sheet['E8'].value] == grid['period_labels']
    assert sheet['D9'].value == '(unaudited)' and sheet['E9'].value is None
    assert [c.value for c in sheet[11][2:]] == ['Cash', 1000, 900]
    assert sheet['C12'].value == grid['rows'][2][0]
    assert all(c.font.bold and c.border.top.style for c in sheet[13][2:])
    assert sheet['C14'].value == grid['source_note']
    book.close()


def test_ambiguous_or_multiple_balance_pages_keep_originals_with_warning():
    plan = StatementPagePlan('unused.pdf', 'manual', (1,), (1,), (1,),
                             (StatementPage(1, 'data', 'BalanceSheet', statement_types=('BalanceSheet', 'IncomeStatement')),))
    original = [dict(source_page=1, statement_type='BalanceSheet')]
    tables, issues = recover_pdf_balance_tables('unused.pdf', original, {}, plan)
    assert tables is original
    assert issues[0]['Status'] == 'WARN' and 'ambiguous' in issues[0]['Detail']


def test_missing_amount_is_not_shifted_or_filled_with_zero():
    page, _, extracted, words = example()
    extracted[3][4] = ''
    words[:] = [w for w in words if w['text'] != '90']
    with pytest.raises(BalanceLayoutError, match='missing date cell'):
        balance_period_grid(page, page.find_tables())


def row_wide_example():
    """Shading defines rows but no amount cells; date rules define columns."""
    entries = [
        ('ASSETS:', '', ''),
        ('Current assets:', '', ''),
        ('Cash', '$ 1,100', '$ 990'),
        ('Total assets', '$ 1,100', '$ 990'),
        ('LIABILITIES AND SHAREHOLDERS’ EQUITY:', '', ''),
        ('Accounts payable', '1,000', '900'),
        ('Common stock, $0.001 par, 500,000 shares authorized;\n100 shares outstanding', '100', '90'),
        ('Total liabilities and shareholders’ equity', '$ 1,100', '$ 990'),
    ]
    words = [dict(text=t, x0=x, x1=x+65, top=y, bottom=y+8)
             for x, date, year in [(405, 'June 30,', '2026'), (505, 'December 31,', '2025')]
             for t, y in [(date, 78), (year, 88)]]
    physical, values = [], []
    for i, (label, left, right) in enumerate(entries):
        top = 100 + 26 * i
        cells = [(20, top, 580, top+26)]
        texts = [' '.join(v for v in (label, left, right) if v)]
        if i == 0:
            cells = [(20, top, 400, top+26), (400, top, 580, top+26)]
            texts = [label, '']
        physical.append(SimpleNamespace(cells=cells))
        values.append(texts)
        for line_index, text in enumerate(label.splitlines()):
            words.append(dict(text=text, x0=25, x1=390 if i != 4 else 570,
                              top=top+2+line_index*10, bottom=top+10+line_index*10))
        for x, text in [(405, left), (505, right)]:
            if text:
                words.append(dict(text=text, x0=x, x1=x+65, top=top+12, bottom=top+20))
    table = SimpleNamespace(rows=physical, bbox=(20, 100, 580, 308), extract=lambda: deepcopy(values))
    text = '\n'.join(['Example Company', 'Balance Sheets', 'June 30, 2026 December 31, 2025'] +
                     [' '.join(v for v in entry if v) for entry in entries] +
                     ['See accompanying Notes to Consolidated Financial Statements.'])
    edges = [dict(x0=x, x1=x+80, top=100, bottom=100) for x in (400, 500)]
    page = SimpleNamespace(width=612, edges=edges, find_tables=lambda: [table], extract_text=lambda: text,
                           extract_words=lambda **kwargs: deepcopy(words))
    return page, table, values, words


def test_row_wide_cells_recover_both_sides_and_keep_share_counts_in_label():
    page, _, _, _ = row_wide_example()
    grid = balance_period_grid(page, page.find_tables())
    assert grid['period_labels'] == ['June 30, 2026', 'December 31, 2025']
    assert len(grid['rows']) == 8 and grid['row_kinds'].count('data') == 5
    assert grid['rows'][2] == ['Cash', '$ 1,100', '$ 990']
    assert grid['rows'][4] == ['LIABILITIES AND SHAREHOLDERS’ EQUITY:', '', '']
    assert grid['rows'][6] == [
        'Common stock, $0.001 par, 500,000 shares authorized; 100 shares outstanding', '100', '90']
    assert grid['rows'][-1] == ['Total liabilities and shareholders’ equity', '$ 1,100', '$ 990']
    assert all(grid['row_sources']) and grid['spans'] == []


@pytest.mark.parametrize('mutation, message', [
    ('date', 'source date columns'), ('crossing', 'crosses date columns'),
    ('missing', 'missing date cell'), ('duplicate', 'ambiguous numeric text'),
    ('value', 'source words'), ('truncated', 'complete source body')])
def test_row_wide_recovery_rejects_ambiguous_or_incomplete_source(mutation, message):
    page, table, values, words = row_wide_example()
    if mutation == 'date':
        words[:] = [w for w in words if w['text'] != '2026']
    elif mutation == 'crossing':
        next(w for w in words if w['text'] == '100')['x1'] = 510
    elif mutation == 'missing':
        words[:] = [w for w in words if w['text'] != '90']
        values[6][0] = values[6][0].removesuffix(' 90')
    elif mutation == 'duplicate':
        duplicate = deepcopy(next(w for w in words if w['text'] == '100'))
        words.append(duplicate)
        values[6][0] += ' 100'
    elif mutation == 'value':
        values[6][0] = values[6][0].replace('100 90', '101 90')
    else:
        table.rows.pop()
        values.pop()
    with pytest.raises(BalanceLayoutError, match=message):
        balance_period_grid(page, page.find_tables())
